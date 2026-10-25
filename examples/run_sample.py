"""
run_sample.py — Run the whole pipeline on the synthetic sample case and check it against the known truth.

Use this to confirm the pipeline works in a new environment (a new computer, operating system or Python version).
It needs no patient data and no one to click: the sample case is synthetic (see make_sample_case.py) and
auto_mark.py stands in for the operator. Each stage is a real command-line script, run as its own process:

    0  make_sample_case.py            builds the synthetic video, force log, tracking sheet and calibration files
    1  index_video.py                 frame table from the video's stored timestamps
    2  parse_force_api.py             force log -> Excel/CSV
    3  extract_point_forces.py        force samples in each window plus the buffer, with roles
    4  plan_measurements.py --yes     summary of force values and frames, default stride
    5  export_calibration_frame.py    the ruler frame and each point's calibration frame (the first frame of its buffer)
    6  examples/auto_mark.py (x4)     camera check, then marks every planned frame the check does not flag (stands in
                                      for calibrate_reference.py and mark_suture_entry.py, which need a person)
    7  mark_suture_entry.py --verify  every marked frame is re-decoded and checked, and checked against the camera again
    8  merge_distance_force.py        the per-frame table: one row per planned frame, its distance and its force

and then checks the results against the truth built into the sample (frame rate, selected force rows, every
measured distance, the stored timestamps, the force paired with every frame, and that the camera check found the
zoom built into the video). The exit code is 0 only if every check passes.

Usage:
    python3 examples/run_sample.py [case_id] [--fresh] [--quiet]
        --fresh   delete and rebuild the sample case first (only ever deletes a folder that holds a sample)
"""

import argparse
import json
import shutil
import subprocess
import sys
import time
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import make_sample_case as mk  # noqa: E402

TOLERANCE_MM_FRAME = 0.05        # per measured frame, against the true gap at that frame's stored time
TOLERANCE_MM_PAIRING = 0.06      # fitted distance at the force second, away from the gap's kinks


def stage(number, label, cmd, quiet):
    start = time.time()
    done = subprocess.run([sys.executable, *[str(c) for c in cmd]], cwd=ROOT, capture_output=True, text=True)
    elapsed = time.time() - start
    if done.returncode != 0:
        print(f"[{number}] {label}: FAILED after {elapsed:.1f} s\n{done.stdout[-1500:]}\n{done.stderr[-1500:]}")
        sys.exit(1)
    if not quiet:
        print(f"[{number}] {label}  ({elapsed:.1f} s)")
    return done.stdout


def sheet_rows(path, name):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[name]
    head = [c.value for c in ws[1]]
    rows = [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    wb.close()
    return rows


class Report:
    def __init__(self):
        self.results = []

    def check(self, name, ok, detail=""):
        self.results.append((name, bool(ok), detail))
        print(f"  {'PASS' if ok else 'FAIL'}  {name}" + (f"  ({detail})" if detail else ""))

    @property
    def ok(self):
        return all(ok for _, ok, _ in self.results)


def run(case_id="SAMPLE001", fresh=False, quiet=False):
    case_dir = ROOT / "Videos" / case_id
    fd = case_dir / "force_data"
    truth_path = case_dir / "ground_truth.json"
    if case_dir.exists() and not truth_path.exists():
        sys.exit(f"ERROR: {case_dir} exists but is not a sample case (no ground_truth.json). "
                 "Choose another case_id; this script will not touch it.")
    if fresh and case_dir.exists():
        shutil.rmtree(case_dir)
    if not case_dir.exists():
        if not quiet:
            print("[0] make_sample_case.py")
        mk.build(case_id, ROOT, quiet=True)
    truth = json.loads(truth_path.read_text())
    video = case_dir / truth["video"]["file"]
    sheet = case_dir / "sample_tracking_sheet.xlsx"
    knots = [tuple(k) for k in truth["gap_knots_s_mm"]]

    stage(1, "index_video.py", ["index_video.py", video, case_id], quiet)
    stage(2, "parse_force_api.py", ["parse_force_api.py", fd / f"{case_id}_force_raw.json", case_id], quiet)
    stage(3, "extract_point_forces.py", ["extract_point_forces.py", case_id, "--tracking-sheet", sheet], quiet)
    plan_out = stage(4, "plan_measurements.py --yes", ["plan_measurements.py", case_id, "--yes"], quiet)
    stage(5, "export_calibration_frame.py (ruler)", ["export_calibration_frame.py", video, case_id, "ruler", "--time",
                                                      f"{truth['ruler_time_s']}"], quiet)
    for p in "ABCD":
        stage(5, f"export_calibration_frame.py (point {p})", ["export_calibration_frame.py", video, case_id, p], quiet)
    for p in "ABCD":
        stage(6, f"auto_mark.py (point {p})", ["examples/auto_mark.py", video, case_id, p], quiet)
    for p in "ABCD":
        stage(7, f"mark_suture_entry.py --verify (point {p})", ["mark_suture_entry.py", video, case_id, p, "--verify"], quiet)
    stage(8, "merge_distance_force.py", ["merge_distance_force.py", case_id], quiet)

    print("\nChecking the results against the known truth")
    rep = Report()

    meta = json.loads((fd / f"{case_id}_frame_times.json").read_text())
    rep.check("frame table has every frame of the video", meta["frames"] == truth["video"]["frames"],
              f"{meta['frames']} frames")
    rep.check("frame rate read from the stored timestamps is 29.97, not 30",
              abs(meta["mean_fps"] - truth["video"]["fps"]) < 1e-3, f"{meta['mean_fps']:.5f} fps")

    force = json.loads((fd / f"{case_id}_force_raw.json").read_text())
    present = {a: {e["x"] for e in force[a]} for a in ("USM1", "USM3")}
    seconds = present["USM1"] | present["USM3"]
    sel = {r["Point"]: r for r in sheet_rows(fd / f"{case_id}_points_ABCD_force.xlsx", "Selection") if r["Point"] in list("ABCD")}
    ok = True
    for p, s in truth["window_starts_s"].items():
        want = len([x for x in seconds if s <= x <= s + truth["window_length_s"]])
        ok &= sel[p]["Window rows"] == want and sel[p]["Buffer-before rows"] > 0 and sel[p]["Buffer-after rows"] > 0
    rep.check("force samples selected per window, with buffer rows on both sides", ok)

    plan = json.loads((fd / f"{case_id}_measurement_plan.json").read_text())
    planned = sum(v["planned_frames"] for v in plan["points"].values())
    marks = json.loads((fd / f"{case_id}_frame_marks.json").read_text())
    records = [r for pt in marks["points"].values() for r in pt["frames"].values()]
    n_marked = sum(1 for r in records if not r.get("skipped"))
    camera_skips = [r for r in records if r.get("skipped")]
    rep.check("every planned frame was marked, or skipped because of the camera",
              n_marked + len(camera_skips) == planned and all(str(r.get("skip_reason", "")).startswith("camera")
                                                              for r in camera_skips),
              f"{n_marked} marked + {len(camera_skips)} skipped for the camera = {planned} planned frames")

    frames_dir = case_dir / "calibration_frames"
    ok = True
    for p, k in truth["calibration_frames"].items():
        sidecar = json.loads((frames_dir / f"{case_id}_{p}_frame{k}.json").read_text())
        ok &= sidecar["frame"] == k and sidecar["is_default_frame"] and sidecar["position"] == "buffer-before"
    rep.check("each point's calibration frame is the first frame of its buffer, exported exactly",
              ok, f"frames {sorted(truth['calibration_frames'].values())}")

    import csv
    with open(fd / f"{case_id}_frame_distances.csv", newline="") as f:
        table = list(csv.DictReader(f))
    frames = [r for r in table if r["status"] == "marked"]
    # the true gap as the camera shows it: the gap times the zoom at that frame's stored time (1.0 when the view is steady;
    # a frame at 1.2% zoom, below the 5% flag, really is measured 1.2% too long, and that is not a measurement error)
    errs = np.array([float(r["distance_mm"]) - mk.gap_mm(float(r["time_s"]), knots) * mk.zoom_at(float(r["time_s"]))
                     for r in frames])
    rep.check("every measured distance matches the true gap at that frame's stored time",
              np.abs(errs).max() <= TOLERANCE_MM_FRAME,
              f"max |error| {np.abs(errs).max():.3f} mm over {len(errs)} measurements")

    # the camera check: the zoom built into the video must be found, and nothing steady may be flagged
    zoom = lambda r: mk.zoom_at(float(r["time_s"]))             # noqa: E731  (the true zoom at the frame's stored time)
    scale = lambda r: float(r["scale_vs_calibration"]) if r["scale_vs_calibration"] != "" else None   # noqa: E731
    must_flag = [r for r in table if abs(zoom(r) - 1) > 0.08]
    must_not = [r for r in table if abs(zoom(r) - 1) < 0.02]
    flagged_ok = [r for r in must_flag if scale(r) is not None and abs(scale(r) - 1) > 0.05
                  and r["status"] == "skipped" and r["skip_reason"].startswith("camera")]
    steady_ok = [r for r in must_not if r["status"] == "marked" and scale(r) is not None and abs(scale(r) - 1) <= 0.05]
    rep.check("the camera check found the zoom in the video: every clearly zoomed frame is flagged and skipped for the camera",
              len(must_flag) >= 6 and len(flagged_ok) == len(must_flag),
              f"{len(flagged_ok)} of {len(must_flag)} zoomed frames")
    rep.check("the camera check left every steady frame alone: marked, scale within 5% of the calibration frame",
              len(steady_ok) == len(must_not), f"{len(steady_ok)} of {len(must_not)} steady frames")
    errs_scale = [abs(scale(r) - zoom(r)) for r in table if scale(r) is not None and r["status"] != "not marked"]
    rep.check("the estimated scale matches the true zoom (within 4%)", errs_scale and max(errs_scale) < 0.04,
              f"max |error| {max(errs_scale):.3f} over {len(errs_scale)} frames")
    with_cal = [r for r in frames if r["calibration_frame"] != ""]
    rep.check("every marked row records the calibration frame its px/mm came from",
              len(with_cal) == len(frames) and all(int(r["calibration_frame"]) == truth["calibration_frames"][r["point"]]
                                                  for r in with_cal), f"{len(with_cal)} rows")

    sheet_frames = sheet_rows(fd / f"{case_id}_points_ABCD_force_merged.xlsx", "Frames")
    rep.check("the Frames sheet and the CSV hold the same table",
              len(sheet_frames) == len(table) and [str(r["frame"]) for r in sheet_frames] == [r["frame"] for r in table],
              f"{len(table)} rows")
    rep.check("one row per planned frame, none of them listed twice",
              len(table) == planned and len({(r["point"], r["frame"]) for r in table}) == len(table),
              f"{len(table)} rows for {planned} planned frames")

    pts_index = {int(r["frame"]): int(r["pts"]) for r in csv.DictReader(open(fd / f"{case_id}_frame_times.csv", newline=""))}
    tb_num, tb_den = (int(v) for v in meta["time_base"].split("/"))
    first_pts = min(pts_index.values())
    bad_pts = [r for r in table if int(r["pts"]) != pts_index[int(r["frame"])]
               or abs(float(r["time_s"]) - (int(r["pts"]) - first_pts) * tb_num / tb_den) > 1e-6]
    rep.check("every row's pts is the stored timestamp, and time_s recomputes from pts and the time base",
              not bad_pts, f"time base {tb_num}/{tb_den}, first pts {first_pts}")

    bad_force = []
    for r in table:
        x = float(r["force_second"])
        for arm in ("USM1", "USM3"):
            want = next((e["y"] for e in force[arm] if e["x"] == x), None)        # None: the log has no sample there
            got = float(r[arm]) if r[arm] != "" else None
            if (got is None) != (want is None) or (want is not None and abs(got - want) > 1e-6):
                bad_force.append((r["point"], r["frame"], arm))
        total = sum(next((e["y"] for e in force[arm] if e["x"] == x), 0.0) for arm in ("USM1", "USM3"))
        if abs(float(r["Total Force (N)"]) - total) > 1e-6:
            bad_force.append((r["point"], r["frame"], "Total"))
    rep.check("every frame's force columns equal the force log at its force second (blank where the log has no "
              "sample; the total counts a blank as 0 N)", not bad_force, f"{len(table)} rows")

    knot_times = np.array([t for t, _ in knots])
    worst = 0.0
    n_pair = 0
    for p in "ABCD":
        for r in sheet_rows(fd / f"{case_id}_points_ABCD_force_merged.xlsx", f"Summary {p} (from Frames)"):
            if r.get("Fit at x (mm)") is None or r["N frames"] is None or r["N frames"] < 3:
                continue
            x = r["Seconds"]
            if np.min(np.abs(knot_times - x)) < 1.0:
                continue
            worst = max(worst, abs(r["Fit at x (mm)"] - mk.gap_mm(x, knots)))
            n_pair += 1
    rep.check("distance at each force second matches the true gap at that second (force-video pairing)",
              n_pair > 20 and worst <= TOLERANCE_MM_PAIRING, f"max |error| {worst:.3f} mm over {n_pair} force seconds")

    # why stored timestamps matter: what the old frame = int(x * 30) rule would have picked
    x = float(truth["window_starts_s"]["D"] + 5)
    naive = int(x * 30)
    t_naive = naive * 1001 / 30000.0
    print(f"\nFor interest: for force second {x:g}, the rule frame = int(x * 30) picks frame {naive}, which is at "
          f"{t_naive:.3f} s on the video clock ({(t_naive - x) * 1000:+.0f} ms off). The true gap there differs from the "
          f"gap at {x:g} s by {abs(mk.gap_mm(t_naive, knots) - mk.gap_mm(x, knots)):.2f} mm. "
          "This pipeline finds frames by their stored timestamps instead.")
    print("\n" + ("ALL CHECKS PASSED" if rep.ok else "SOME CHECKS FAILED"))
    return 0 if rep.ok else 1


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Run the pipeline on the synthetic sample case and check the results.")
    ap.add_argument("case_id", nargs="?", default="SAMPLE001")
    ap.add_argument("--fresh", action="store_true", help="rebuild the sample case first")
    ap.add_argument("--quiet", action="store_true")
    a = ap.parse_args()
    sys.exit(run(a.case_id, a.fresh, a.quiet))
