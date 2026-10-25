"""
merge_distance_force.py — Pair every measured frame's crural distance with the force of its force second.

This is the last stage of the pipeline. Its product is ONE TABLE WITH ONE ROW PER MEASURED FRAME: the frame's
distance and the force paired with it. Strain, moduli, curves and every other calculation are done by the analyst
from this table (see the methods paper); the pipeline does not fit anything.

Reads (never modified):
    Videos/<case_id>/force_data/<case_id>_points_ABCD_force.xlsx  (from extract_point_forces.py)
    Videos/<case_id>/force_data/<case_id>_frame_marks.json        (from mark_suture_entry.py)
    Videos/<case_id>/force_data/<case_id>_frame_times.csv         (frame timestamps, from index_video.py)

Writes (new files):
    Videos/<case_id>/force_data/<case_id>_frame_distances.csv           the per-frame table
    Videos/<case_id>/force_data/<case_id>_points_ABCD_force_merged.xlsx  the same table as sheet "Frames", plus
        "Summary <point> (from Frames)"   one row per force second (secondary, derived from the Frames sheet)
        "Selection"                       what extract_point_forces.py selected
        "Merge info"                      version, time base, counts, how to read the columns

The per-frame table has one row per planned frame of every selected force second (window and buffer), in the order
point, force second, frame. Columns:
    case_id, point                  case ID and suture point (A to D)
    role                            window | buffer-before | buffer-after: where the force second lies relative to the noted window
    force_second, force_time_hms    the force sample's time in the force log (seconds, and HH:MM:SS)
    frame                           raw frame number in the video (0 = first frame, in presentation order)
    pts                             the frame's STORED presentation timestamp, in time-base ticks (an integer, as read from the file)
    time_s                          (pts - first frame's pts) x time base: the video time a player shows, in seconds
    time_hms                        the same as HH:MM:SS.mmm
    offset_frames, offset_ms        the frame's distance from the force second's centre frame, in frames and in ms
                                    (frame time minus the force sample's target time)
    status                          marked | skipped | not marked
    skip_reason                     why a skipped frame was skipped (for example "camera: scale 0.76 vs calibration"); empty otherwise
    distance_mm, distance_px        the crural distance in millimetres and pixels (empty unless marked)
    left_x, left_y, right_x, right_y  the two clicks, in pixels (left crus first)
    px_per_mm                       the scale the distance was converted with (empty unless marked)
    calibration_frame               the raw frame number of the calibration frame that scale came from (empty if unknown)
    scale_vs_calibration            the camera check: the scene's scale on this frame relative to the calibration frame
                                    (1.00 = same zoom; empty if no check was made or no estimate was possible)
    USM1 ... USM4                   the force of each logged arm in N, as in the force log; BLANK where the log has no sample
    Total Force (N)                 the sum of the arms, where a blank arm counts as 0 N

The "Summary ... (from Frames)" sheets are derived from the Frames sheet: for each force second, the mean, SD, N,
min, max and median of its marked frames' distances, and a straight line through distance against frame time evaluated at
the sample time. The frames of one force second lie up to +/-0.47 s apart (14 frames at 30 fps), and the gap can change by
several mm per second, so the SD mixes real tissue motion, click noise and timing ambiguity. It is not pure measurement error.

The frames that belong to each force second come from the same parameters the marks were made with (stored in the marks
file), or from the options below to re-summarise the same frames under another stride, e.g. the sparse subset of a denser
pass. A half-window of 15 frames or more is refused: every frame must belong to one force second.

Usage:
    python3 merge_distance_force.py <case_id> [--half-window-frames N] [--stride N]
            [--time-offset-s S] [--window-center-shift-s S]
"""

import argparse
import csv
import os
import sys
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.styles import Font

import camera_check as cc
import framelib as fl

BASE = Path(__file__).resolve().parent
POINTS = ["A", "B", "C", "D"]
ID_COLUMNS = [
    "case_id", "point", "role", "force_second", "force_time_hms", "frame", "pts", "time_s", "time_hms",
    "offset_frames", "offset_ms", "status", "skip_reason", "distance_mm", "distance_px",
    "left_x", "left_y", "right_x", "right_y", "px_per_mm", "calibration_frame", "scale_vs_calibration",
]
TOTAL = "Total Force (N)"
SUMMARY_COLS = [
    TOTAL, "Distance (mm)", "Distance SD (mm)", "N frames", "Frames planned",
    "Distance min (mm)", "Distance max (mm)", "Distance median (mm)", "Distance (px)",
    "Fit at x (mm)", "Slope (mm/s)", "Fit resid SD (mm)",
]
NOT_ARM = ("Timestamp", "Seconds", "In Window", "Role")


def summarise(recs, t_rel):
    """Statistics of the measured frames of one force second. recs: non-skipped records."""
    n = len(recs)
    out = {"n": n}
    if n == 0:
        return out
    d = np.array([r["distance_mm"] for r in recs], float)
    out.update(mean=float(d.mean()), min=float(d.min()), max=float(d.max()), median=float(np.median(d)),
               px=float(np.mean([r["distance_px"] for r in recs])))
    if n >= 2:
        out["sd"] = float(d.std(ddof=1))
        t = np.asarray(t_rel, float)
        if np.ptp(t) > 0:
            slope, intercept = np.polyfit(t, d, 1)
            out["fit_at_x"], out["slope"] = float(intercept), float(slope)
            if n >= 3:
                out["fit_sd"] = float(np.sqrt(np.sum((d - (slope * t + intercept)) ** 2) / (n - 2)))
    return out


def r(v, nd=4):
    return None if v is None else round(v, nd)


def force_time_hms(x):
    """HH:MM:SS of a force second (with milliseconds only if the second is not a whole number)."""
    s = fl.fmt_time(x)
    return s[:8] if s.endswith(".000") else s


def parse_args(argv=None):
    ap = argparse.ArgumentParser(description="Pair every measured frame's distance with its force: the per-frame table.")
    ap.add_argument("case_id")
    ap.add_argument("--half-window-frames", dest="half_window_frames", type=int)
    ap.add_argument("--stride", type=int)
    ap.add_argument("--time-offset-s", dest="time_offset_s", type=float)
    ap.add_argument("--window-center-shift-s", dest="window_center_shift_s", type=float)
    ap.add_argument("--version", action="version", version=f"%(prog)s {fl.PIPELINE_VERSION}")
    return ap.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    case_id = args.case_id
    force_dir = BASE / "Videos" / case_id / "force_data"
    force_path = force_dir / f"{case_id}_points_ABCD_force.xlsx"
    marks_path = force_dir / f"{case_id}_frame_marks.json"
    out_xlsx = force_dir / f"{case_id}_points_ABCD_force_merged.xlsx"
    out_csv = force_dir / f"{case_id}_frame_distances.csv"

    for p, hint in ((force_path, "run extract_point_forces.py first"),
                    (marks_path, "run mark_suture_entry.py first")):
        if not p.exists():
            sys.exit(f"ERROR: {p} not found - {hint}")

    index = fl.load_index(force_dir, case_id)
    store = fl.MarksStore(marks_path, case_id)
    vid = store.data.get("video", {})
    if vid.get("frames") not in (None, index.n):
        sys.exit("ERROR: the frame table and the marks come from videos with different frame counts")

    wb = openpyxl.load_workbook(force_path)
    bold = Font(bold=True)
    frame_rows, info, arms = [], [], []

    for label in POINTS:
        sheet = f"Point {label}"
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        headers = [c.value for c in ws[1]]
        col = {h: i for i, h in enumerate(headers)}
        arm_cols = [h for h in headers if h not in NOT_ARM]
        arms += [a for a in arm_cols if a not in arms]
        for h in SUMMARY_COLS:
            c = ws.max_column + 1
            ws.cell(row=1, column=c, value=h).font = bold
            col[h] = c - 1
            ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 16

        given = {k: getattr(args, k) for k in fl.PARAM_DEFAULTS}
        params = fl.resolve_params(index, **{**(store.last_params(label) or {}),
                                             **{k: v for k, v in given.items() if v is not None}})
        try:
            fl.check_half_window(index, params["half_window_frames"])
        except ValueError as e:
            sys.exit(f"ERROR: point {label}: {e}")
        check = store.get_scale_check(label) or {}
        planned = measured = skipped = unmarked = camera_skips = camera_flags = 0
        with_data = {"window": 0, "buffer": 0}
        for row in ws.iter_rows(min_row=2):
            secs = row[col["Seconds"]].value
            forces = {a: row[col[a]].value for a in arm_cols}
            total = sum((v or 0) for v in forces.values())
            row[col[TOTAL]].value = total
            if secs is None:
                continue
            x = round(float(secs), 3)
            if "Role" in col:
                role = row[col["Role"]].value
            else:   # workbook made by an older extraction: only its In Window rows were measured
                role = "window" if row[col["In Window"]].value == "YES" else None
            to_measure = role in fl.ROLES
            centre, frames = fl.sample_frames(index, x, **params)
            recs, t_rel = [], []
            t_target = fl.target_time(x, params)
            for sf in frames:
                rec = store.get(label, sf.frame)
                status = "not marked" if rec is None else ("skipped" if rec.get("skipped") else "marked")
                est = store.get_scale(label, sf.frame)
                if to_measure:
                    planned += 1
                    measured += status == "marked"
                    skipped += status == "skipped"
                    unmarked += status == "not marked"
                    if status == "skipped" and str(rec.get("skip_reason") or "").startswith("camera"):
                        camera_skips += 1
                    if status == "marked" and check and cc.frame_flag(est, check.get("threshold", cc.DEFAULT_THRESHOLD),
                                                                      check.get("min_inliers", cc.DEFAULT_MIN_INLIERS)):
                        camera_flags += 1
                    marked = status == "marked"
                    t = float(index.pts_s[sf.frame])
                    frame_rows.append({
                        "case_id": case_id, "point": label, "role": role, "force_second": x,
                        "force_time_hms": force_time_hms(x), "frame": sf.frame, "pts": int(index.pts[sf.frame]),
                        "time_s": r(t, 6), "time_hms": fl.fmt_time(t), "offset_frames": sf.offset_frames,
                        "offset_ms": r(sf.offset_ms, 2), "status": status,
                        "skip_reason": rec.get("skip_reason") if status == "skipped" else None,
                        "distance_mm": r(rec.get("distance_mm")) if marked else None,
                        "distance_px": r(rec.get("distance_px"), 2) if marked else None,
                        "left_x": rec["left_crus_px"][0] if marked else None,
                        "left_y": rec["left_crus_px"][1] if marked else None,
                        "right_x": rec["right_crus_px"][0] if marked else None,
                        "right_y": rec["right_crus_px"][1] if marked else None,
                        "px_per_mm": rec.get("px_per_mm") if marked else None,
                        "calibration_frame": rec.get("calibration_frame") if marked else None,
                        "scale_vs_calibration": est.get("scale") if est else None,
                        **forces, TOTAL: total})
                if status == "marked":
                    recs.append(rec)
                    t_rel.append(float(index.pts_s[sf.frame]) - t_target)
            st = summarise(recs, t_rel)
            row[col["Frames planned"]].value = len(frames) if to_measure else None
            row[col["N frames"]].value = st["n"] if st["n"] else None
            if st["n"]:
                with_data["window" if role == "window" else "buffer"] += 1
                row[col["Distance (mm)"]].value = r(st["mean"])
                row[col["Distance SD (mm)"]].value = r(st.get("sd"))
                row[col["Distance min (mm)"]].value = r(st["min"])
                row[col["Distance max (mm)"]].value = r(st["max"])
                row[col["Distance median (mm)"]].value = r(st["median"])
                row[col["Distance (px)"]].value = r(st["px"], 2)
                row[col["Fit at x (mm)"]].value = r(st.get("fit_at_x"))
                row[col["Slope (mm/s)"]].value = r(st.get("slope"))
                row[col["Fit resid SD (mm)"]].value = r(st.get("fit_sd"))
        ws.title = f"Summary {label} (from Frames)"
        print(f"  Point {label}: {measured} frames marked, {skipped} skipped, {unmarked} not marked "
              f"(of {planned} planned); distance summarised for {with_data['window']} window and "
              f"{with_data['buffer']} buffer force seconds")
        info.append([label, *(params[k] for k in fl.PARAM_DEFAULTS), with_data["window"], with_data["buffer"],
                     planned, measured, skipped, unmarked, camera_skips,
                     camera_flags if check else None])

    columns = ID_COLUMNS + arms + [TOTAL]
    frame_rows.sort(key=lambda d: (d["point"], d["force_second"], d["frame"]))
    table = [[d.get(c) for c in columns] for d in frame_rows]

    wf = wb.create_sheet("Frames", 0)
    wf.append(columns)
    for c in wf[1]:
        c.font = bold
    for line in table:
        wf.append(line)
    wf.freeze_panes = "A2"

    ws = wb.create_sheet("Merge info")
    ws.append([f"Pipeline version {fl.PIPELINE_VERSION}. Per-point parameters and frame counts (frames counted over the "
               "force seconds that are measured; every frame belongs to one force second)"])
    ws.append(["Point", *fl.PARAM_DEFAULTS.keys(), "Window seconds with data", "Buffer seconds with data",
               "Frames planned", "Marked", "Skipped", "Not marked", "Skipped for camera",
               "Marked frames flagged by the camera check"])
    for line in info:
        ws.append(line)
    ws.append([])
    ws.append([f"Video: {vid.get('name', index.source_name)}, {index.n} frames. Time base {index.tb_num}/{index.tb_den} s per tick; "
               f"first frame pts {int(index.pts[0])}. time_s = (pts - first frame pts) x time base, so video time zero is the first frame; "
               "time_offset_s = video time minus force time."])
    ws.append(["The Frames sheet (and <case_id>_frame_distances.csv, the same table) is the product: one row per measured frame."])
    ws.append(["The Summary sheets are derived from the Frames sheet: for each force second the mean, SD, N, min, max and median of "
               "the marked frames' distances. The SD mixes tissue motion, click noise and timing ambiguity; it is not pure "
               "measurement error."])
    ws.append(["Role column: window = between the noted start and end; buffer-before / buffer-after = in the buffer. "
               "Buffer rows are measured and kept; filter the role column to leave them out."])
    ws.append([f"Force: one column per logged arm (USM1 to USM4) as in the force log, blank where the log has no sample; "
               f"{TOTAL} adds the arms and counts a blank arm as 0 N."])
    ws.append(["scale_vs_calibration: the camera check (camera_check.py): the scene's scale on the frame relative to the point's "
               "calibration frame. 1.00 = same zoom. A scale far from 1 means the pixel-to-millimetre scale of the point does not "
               "apply to that frame."])
    for c in ws[2]:
        c.font = bold

    tmp = out_xlsx.with_name(out_xlsx.name + ".tmp")
    wb.save(tmp)
    os.replace(tmp, out_xlsx)
    with open(out_csv, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(columns)
        w.writerows(table)
    print(f"Saved: {out_xlsx}")
    print(f"Saved: {out_csv}")


if __name__ == "__main__":
    main()
