"""
calibrate_reference.py — Establish pixel/mm calibration using anatomical landmarks.

Two modes:

  FULL MODE (calibration frame with ruler):
    Step 1: Click two ruler tick marks → type known mm → get px/mm
    Step 2: Click two blood vessel landmarks → records distance in mm (ground truth)
    Usage: python3 calibrate_reference.py <image> <case_id>
    This mode also writes the calibration columns (px_per_mm_ruler, landmark_mm, ...) into the row of
    the case in the tracking sheet IN PLACE, with no backup; the calibration JSON holds the same
    results, and nothing later reads those columns.

  LANDMARK-ONLY MODE (the calibration frame of each point A-D, no ruler):
    Just click the two blood vessels. Uses the known landmark_mm from
    <case_id>_calibration.json to compute px/mm for this frame's depth. A small panel on the right shows
    the previous reference frame with its two landmark marks (point A shows the ruler frame, B shows A's,
    C shows B's, D shows C's; the ruler frame when the previous point is not calibrated yet), so the same
    two vessels are clicked in the same order.
    Usage: python3 calibrate_reference.py <image> <case_id> <point>

The calibration frame of a point is the FIRST frame of that point's segment (the window start in the
tracking sheet): later frames may already show the tissue deformed by pulling. Export it exactly, at the
video's native resolution, with export_calibration_frame.py. If the landmarks are not visible there, or the
camera has moved before the forces are collected, pick another frame with pick_calibration_frame.py. The
frame used is recorded in the calibration JSON (frame_info), so deviations can be found afterwards.

Examples:
    python3 calibrate_reference.py Videos/<case_id>/calibration_frame.png <case_id>
    python3 calibrate_reference.py Videos/<case_id>/calibration_frames/<case_id>_A_frame<N>.png <case_id> A

The plot window uses matplotlib's default interactive backend for your platform. Set the
MPLBACKEND environment variable (for example TkAgg or QtAgg) to choose another.
The command exits with an error unless the calibration was completed and saved ("Done" in the title).
"""

import argparse
import json
import sys
from pathlib import Path

import cv2
import numpy as np
import openpyxl

import framelib as fl

BASE = Path(__file__).resolve().parent
PIPELINE_VERSION = fl.PIPELINE_VERSION
POINTS = "ABCD"


def default_sheet():
    return BASE / "Hiatal tension patients.xlsx"


# ------------------------------------------------------------ click logic -----

class Calibration:
    """The click state of one calibration: ruler first (full mode), then two landmarks.

    phase: ruler -> ruler_input (type the ruler length) -> landmark -> done.
    Clicks are pixel coordinates. Two clicks on the same pixel are not a distance, so the second is dropped.
    """

    def __init__(self, landmark_mm=None):
        self.landmark_mm = landmark_mm            # known landmark distance: landmark-only mode
        self.phase = "landmark" if landmark_mm is not None else "ruler"
        self.ruler_clicks, self.landmark_clicks = [], []
        self.ruler_px = self.ruler_mm = self.px_per_mm = None
        self.result = None

    def click(self, x, y):
        point = (int(x), int(y))
        if self.phase == "ruler":
            if self.ruler_clicks and point == self.ruler_clicks[0]:
                return
            self.ruler_clicks.append(point)
            if len(self.ruler_clicks) == 2:
                (x1, y1), (x2, y2) = self.ruler_clicks
                self.ruler_px = float(np.hypot(x2 - x1, y2 - y1))
                self.phase = "ruler_input"
        elif self.phase == "landmark":
            if self.landmark_clicks and point == self.landmark_clicks[0]:
                return
            self.landmark_clicks.append(point)
            if len(self.landmark_clicks) == 2:
                self._finish_landmarks()

    def submit_ruler_mm(self, text):
        """The typed ruler length in mm. False (and no change) if it is not a positive number."""
        if self.phase != "ruler_input":
            return False
        try:
            mm = float(text)
        except (TypeError, ValueError):
            return False
        if not mm > 0:
            return False
        self.ruler_mm = mm
        self.px_per_mm = self.ruler_px / mm
        self.phase = "landmark"
        return True

    def _finish_landmarks(self):
        p1, p2 = self.landmark_clicks
        px = float(np.hypot(p2[0] - p1[0], p2[1] - p1[1]))
        if self.landmark_mm is not None:
            mm = self.landmark_mm
            self.px_per_mm = px / mm
        else:
            mm = px / self.px_per_mm
        self.result = {"px_per_mm": self.px_per_mm, "landmark_px": px, "landmark_mm": mm,
                       "landmark_coords_p1": p1, "landmark_coords_p2": p2, "ruler_mm": self.ruler_mm,
                       "ruler_coords_p1": self.ruler_clicks[0] if self.ruler_clicks else None,
                       "ruler_coords_p2": self.ruler_clicks[1] if len(self.ruler_clicks) > 1 else None}
        self.phase = "done"


# ------------------------------------------------------------------ saving ----

def read_frame_info(img_path):
    """What export_calibration_frame.py recorded about the still (frame number, video time, whether it is the
    default first frame of the segment), from the .json next to it. None for a still made by hand."""
    side = Path(img_path).with_suffix(".json")
    if not side.exists():
        return None
    try:
        return json.loads(side.read_text())
    except (OSError, ValueError):
        return None


def save_calibration(session, case_id, label, img_path, tracking_sheet=None):
    """Write the calibration JSON (and, in full mode, the tracking-sheet columns). Returns the JSON path."""
    r = session.result
    px_per_mm = r["px_per_mm"]
    landmark_only = session.landmark_mm is not None
    cal = {
        "case": case_id,
        "clip": label if label else "calibration",
        "calibration_frame": str(img_path),
        "frame_info": read_frame_info(img_path),
        "mode": "landmark_only" if landmark_only else "ruler",
        "px_per_mm": round(px_per_mm, 4),
        "mm_per_px": round(1 / px_per_mm, 6),
        "landmark_px": round(r["landmark_px"], 2),
        "landmark_mm": round(r["landmark_mm"], 4),
        "landmark_coords_p1": r["landmark_coords_p1"],
        "landmark_coords_p2": r["landmark_coords_p2"],
        "pipeline_version": PIPELINE_VERSION,
    }
    if not landmark_only:
        cal["ruler_mm"] = r["ruler_mm"]
        cal["ruler_coords_p1"] = r["ruler_coords_p1"]
        cal["ruler_coords_p2"] = r["ruler_coords_p2"]

    suffix = f"_{label}" if label else ""
    fdir = BASE / "Videos" / case_id / "force_data"
    fdir.mkdir(parents=True, exist_ok=True)
    out_json = fdir / f"{case_id}{suffix}_calibration.json"
    with open(out_json, "w") as f:
        json.dump(cal, f, indent=2)
    print(f"Saved: {out_json}")

    if not landmark_only:
        sheet = Path(tracking_sheet) if tracking_sheet else default_sheet()
        if not sheet.exists():
            print(f"Note: tracking sheet not found ({sheet}); the results are in the calibration JSON only.")
        else:
            update_tracking_sheet(sheet, case_id, px_per_mm, r)
            print("Updated the tracking sheet (in place, no backup)")
    return out_json


def update_tracking_sheet(sheet_path, case_id, px_per_mm, r):
    wb = openpyxl.load_workbook(sheet_path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]

    def get_or_add_col(name):
        if name in headers:
            return headers.index(name) + 1
        col = len(headers) + 1
        ws.cell(row=1, column=col, value=name)
        headers.append(name)
        return col

    col_pxmm = get_or_add_col("px_per_mm_ruler")
    col_mmpx = get_or_add_col("mm_per_px_ruler")
    col_lmk_px = get_or_add_col("landmark_px")
    col_lmk_mm = get_or_add_col("landmark_mm")
    col_ruler_mm = get_or_add_col("ruler_distance_mm")

    for row in ws.iter_rows(min_row=2):
        if row[0].value == case_id:
            row[col_pxmm - 1].value = round(px_per_mm, 4)
            row[col_mmpx - 1].value = round(1 / px_per_mm, 6)
            row[col_lmk_px - 1].value = round(r["landmark_px"], 2)
            row[col_lmk_mm - 1].value = round(r["landmark_mm"], 4)
            row[col_ruler_mm - 1].value = r["ruler_mm"]
            break
    wb.save(sheet_path)


# ------------------------------------------------------ the previous frame -----

def resolve_image(path, base=None):
    """An image path as stored in a calibration JSON (absolute, or relative to the repository folder or to
    the working directory), or None if no such file exists."""
    base = BASE if base is None else Path(base)
    p = Path(path)
    for candidate in ([p] if p.is_absolute() else [base / p, Path.cwd() / p]):
        if candidate.exists():
            return str(candidate)
    return None


def previous_reference(case_id, label, base=None):
    """The frame to show beside the one being calibrated: the previous point's (A shows the ruler frame, B shows
    A's, C shows B's, D shows C's), else the ruler frame when the previous point has no calibration yet.

    Returns {name, image (path or None), p1, p2, missing (stored path when the image is gone)} or None."""
    base = BASE if base is None else Path(base)
    fdir = base / "Videos" / case_id / "force_data"
    candidates = []
    if label in POINTS and POINTS.index(label) > 0:
        prev = POINTS[POINTS.index(label) - 1]
        candidates.append((f"point {prev}", fdir / f"{case_id}_{prev}_calibration.json"))
    candidates.append(("the ruler frame", fdir / f"{case_id}_calibration.json"))
    for name, path in candidates:
        if not path.exists():
            continue
        try:
            cal = json.loads(path.read_text())
            p1, p2 = cal["landmark_coords_p1"], cal["landmark_coords_p2"]
            stored = cal["calibration_frame"]
        except (OSError, ValueError, KeyError):
            continue
        image = resolve_image(stored, base)
        return {"name": name, "image": image, "p1": tuple(p1), "p2": tuple(p2),
                "missing": None if image else str(stored)}
    return None


# ----------------------------------------------------------------- window -----

def run_gui(session, img_rgb, case_id, label, img_path, reference, tracking_sheet):
    """Show the frame and take the clicks. Returns True once the calibration is complete and saved."""
    import matplotlib.pyplot as plt
    from matplotlib.widgets import TextBox

    saved = {"done": False}
    landmark_only = session.landmark_mm is not None
    fig = plt.figure(figsize=(13, 10))
    width = 0.78 if reference else 1.0              # the reference panel sits beside the frame, not over it
    ax = fig.add_axes([0, 0.08, width, 0.92])
    ax_box = fig.add_axes([0.3, 0.01, 0.2, 0.05])
    textbox = TextBox(ax_box, "Ruler mm: ", initial="")
    ax_box.set_visible(False)

    ax.imshow(img_rgb)
    ax.axis("off")
    info_text = ax.text(10, img_rgb.shape[0] - 10, "", fontsize=11, color="white",
                        va="bottom", bbox=dict(facecolor="black", alpha=0.6))

    if reference:
        draw_reference(fig, reference)

    if landmark_only:
        ax.set_title(f"Click TWO blood vessel landmarks  |  reference = {session.landmark_mm} mm",
                     fontsize=13, color="white", bbox=dict(boxstyle="round", facecolor="#1a5276", alpha=0.8))
    else:
        ax.set_title("STEP 1: Click TWO ruler tick marks of known distance",
                     fontsize=13, color="white", bbox=dict(boxstyle="round", facecolor="#333333", alpha=0.8))

    def update_info():
        phase = session.phase
        if phase == "ruler":
            msg = f"Ruler clicks: {len(session.ruler_clicks)}/2  — click two tick marks"
        elif phase == "ruler_input":
            msg = "Type the mm distance between the tick marks in the box below, then press Enter"
        elif phase == "landmark":
            n = len(session.landmark_clicks)
            if session.px_per_mm and not landmark_only:
                msg = f"px/mm = {session.px_per_mm:.2f}  |  Landmark clicks: {n}/2  — click two blood vessels"
            else:
                msg = f"Landmark clicks: {n}/2  — click two blood vessels"
        elif phase == "done":
            msg = "Done — saved. You can close the window"
        else:
            msg = ""
        info_text.set_text(msg)
        fig.canvas.draw()

    def onclick(event):
        if event.inaxes != ax or event.xdata is None:
            return
        x, y = int(event.xdata), int(event.ydata)
        before = session.phase, len(session.ruler_clicks), len(session.landmark_clicks)
        session.click(x, y)
        if before == (session.phase, len(session.ruler_clicks), len(session.landmark_clicks)):
            return                                  # nothing changed (ignored click)
        if before[0] == "ruler":
            ax.plot(x, y, "+", color="yellow", markersize=12, markeredgewidth=1.5)
            if session.phase == "ruler_input":
                p1, p2 = session.ruler_clicks
                ax.plot([p1[0], p2[0]], [p1[1], p2[1]], "-", color="yellow", linewidth=1.5)
                ax_box.set_visible(True)
                textbox.set_val("")
            update_info()
        elif before[0] == "landmark":
            ax.plot(x, y, "+", color="cyan", markersize=12, markeredgewidth=1.5)
            if session.phase == "done":
                r = session.result
                p1, p2 = r["landmark_coords_p1"], r["landmark_coords_p2"]
                ax.plot([p1[0], p2[0]], [p1[1], p2[1]], "-", color="cyan", linewidth=1.5)
                ax.annotate(f"Landmark: {r['landmark_px']:.0f} px = {r['landmark_mm']:.2f} mm\n"
                            f"→ {r['px_per_mm']:.2f} px/mm",
                            ((p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2 - 15), color="cyan", fontsize=9,
                            ha="center", bbox=dict(facecolor="black", alpha=0.6))
                print(f"\nLandmark distance: {r['landmark_px']:.1f} px = {r['landmark_mm']:.2f} mm")
                print(f"Calibration: {r['px_per_mm']:.2f} px/mm  |  {1 / r['px_per_mm']:.4f} mm/px")
                save_calibration(session, case_id, label, img_path, tracking_sheet)
                saved["done"] = True
                ax.set_title(f"Done — {r['landmark_px']:.0f} px = {r['landmark_mm']:.2f} mm  |  "
                             f"{r['px_per_mm']:.2f} px/mm", fontsize=13, color="white",
                             bbox=dict(boxstyle="round", facecolor="#1e8449", alpha=0.8))
            update_info()

    def on_ruler_submit(text):
        if not session.submit_ruler_mm(text):
            return
        p1, p2 = session.ruler_clicks
        ax_box.set_visible(False)
        ax.annotate(f"{session.ruler_mm:.1f} mm = {session.ruler_px:.0f} px\n→ {session.px_per_mm:.2f} px/mm",
                    ((p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2 - 15), color="yellow", fontsize=9,
                    ha="center", bbox=dict(facecolor="black", alpha=0.6))
        ax.set_title("STEP 2: Click TWO anatomical landmarks (blood vessels on the crus)",
                     fontsize=13, color="white", bbox=dict(boxstyle="round", facecolor="#1a5276", alpha=0.8))
        update_info()

    textbox.on_submit(on_ruler_submit)
    fig.canvas.mpl_connect("button_press_event", onclick)
    update_info()
    try:
        plt.get_current_fig_manager().full_screen_toggle()
    except Exception:
        pass
    plt.show()
    plt.close(fig)
    return saved["done"]


def draw_reference(fig, reference):
    """The previous reference frame, small, with its two landmark marks numbered in click order."""
    ax_ref = fig.add_axes([0.79, 0.50, 0.20, 0.42])
    ax_ref.axis("off")
    ax_ref.set_title(f"Previous reference: {reference['name']}\nclick the same two vessels, same order",
                     fontsize=9)
    if reference["image"] is None:
        ax_ref.text(0.5, 0.5, f"image not found:\n{reference['missing']}", ha="center", va="center", fontsize=8,
                    transform=ax_ref.transAxes, wrap=True)
        return
    img = cv2.cvtColor(cv2.imread(reference["image"]), cv2.COLOR_BGR2RGB)
    ax_ref.imshow(img)
    (x1, y1), (x2, y2) = reference["p1"], reference["p2"]
    ax_ref.plot([x1, x2], [y1, y2], "-", color="cyan", linewidth=1.2)
    for n, (x, y) in enumerate(((x1, y1), (x2, y2)), 1):
        ax_ref.plot(x, y, "+", color="cyan", markersize=12, markeredgewidth=1.5)
        ax_ref.annotate(str(n), (x, y), color="yellow", fontsize=11, fontweight="bold",
                        xytext=(6, -6), textcoords="offset points")


# ------------------------------------------------------------------- main -----

def parse_args(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("image", help="calibration frame (PNG/TIFF at the video's native resolution)")
    ap.add_argument("case_id")
    ap.add_argument("point", nargs="?", type=str.upper, choices=list(POINTS),
                    help="A-D: landmark-only calibration of that point (needs the ruler calibration first)")
    ap.add_argument("--tracking-sheet", default=None, dest="tracking_sheet",
                    help="spreadsheet that the full calibration writes its columns into [Hiatal tension patients.xlsx]")
    ap.add_argument("--version", action="version", version=f"%(prog)s {PIPELINE_VERSION}")
    return ap.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    case_id, label = args.case_id, args.point

    landmark_mm = None
    if label:
        base_cal = BASE / "Videos" / case_id / "force_data" / f"{case_id}_calibration.json"
        if not base_cal.exists():
            sys.exit(f"ERROR: base calibration file not found: {base_cal}\n"
                     f"       Run the full calibration first: calibrate_reference.py <ruler frame> {case_id}")
        landmark_mm = json.loads(base_cal.read_text())["landmark_mm"]
        print(f"Loaded reference: landmark_mm = {landmark_mm} mm from {base_cal.name}")

    img_bgr = cv2.imread(str(args.image))
    if img_bgr is None:
        sys.exit(f"ERROR: cannot read the image: {args.image}")
    img_rgb = cv2.cvtColor(img_bgr, cv2.COLOR_BGR2RGB)

    reference = previous_reference(case_id, label) if label else None
    session = Calibration(landmark_mm)
    if run_gui(session, img_rgb, case_id, label, args.image, reference, args.tracking_sheet):
        return
    sys.exit("No calibration saved: the window was closed before the calibration was complete "
             "(the title says Done when it is saved). Run the command again and click both landmarks"
             + ("" if label else ", with the ruler first") + ".")


if __name__ == "__main__":
    main()
