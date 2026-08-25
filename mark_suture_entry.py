"""
mark_suture_entry.py — Manually mark suture entry points at 1-second intervals.

Steps through a clip at 1-second intervals (matching force data resolution).
On each frame, click the LEFT crus entry point then the RIGHT crus entry point.
Press SPACE to skip a frame (e.g. if occluded). Press Q to quit early.

Results saved to:
    Videos/<case_id>/force_data/<case_id>_suture_entry.json
    Videos/<case_id>/force_data/<case_id>_suture_entry.xlsx

Usage:
    python3 mark_suture_entry.py <video.mp4> <case_id> <clip_label>

Examples:
    python3 mark_suture_entry.py Videos/YV003/clips/A/YV003_A.mp4 YV003 A
    python3 mark_suture_entry.py Videos/YV003/clips/B/YV003_B.mp4 YV003 B
"""

import sys
import json
import numpy as np
import cv2
import matplotlib
matplotlib.use("MacOSX")
import matplotlib.pyplot as plt
from pathlib import Path
import openpyxl

BASE = Path(__file__).resolve().parent

video_path = sys.argv[1] if len(sys.argv) > 1 else None
case_id    = sys.argv[2] if len(sys.argv) > 2 else "YV003"
clip_label = sys.argv[3] if len(sys.argv) > 3 else None

if not video_path or not clip_label:
    print(__doc__)
    sys.exit(1)

# Load clip calibration
cal_path = BASE / "Videos" / case_id / "force_data" / f"{case_id}_{clip_label}_calibration.json"
if not cal_path.exists():
    print(f"ERROR: Calibration file not found: {cal_path}")
    sys.exit(1)

with open(cal_path) as f:
    cal = json.load(f)
px_per_mm = cal["px_per_mm"]
print(f"Loaded calibration for clip {clip_label}: {px_per_mm:.2f} px/mm")

def ts_to_secs(ts):
    if not ts: return None
    parts = str(ts).split(":")
    return int(parts[0])*3600 + int(parts[1])*60 + float(parts[2])

def secs_to_ts(s):
    h = int(s // 3600)
    m = int((s % 3600) // 60)
    sec = s % 60
    return f"{h:02d}:{m:02d}:{sec:06.3f}"

# Prefer full timestamped video for frame-accurate seeking (abs_sec * fps directly)
full_video = BASE / "Videos" / case_id / f"{case_id}_full_timestamped.mp4"
if full_video.exists():
    seek_video = str(full_video)
    use_full_video = True
    print(f"Using full timestamped video for frame-accurate seeking: {full_video.name}")
else:
    seek_video = video_path
    use_full_video = False
    print(f"Full timestamped video not found — using clip with offset seek")

cap   = cv2.VideoCapture(seek_video)
fps   = cap.get(cv2.CAP_PROP_FPS)
total = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
duration = total / fps

frames = {}  # abs_sec -> (frame_img, absolute_timestamp)

# Load clip start time (used only as fallback)
wb_pat = openpyxl.load_workbook(BASE / "Hiatal tension patients.xlsx")
ws_pat = wb_pat.active
headers = [cell.value for cell in ws_pat[1]]
col = {h: i for i, h in enumerate(headers)}

clip_start_ts = None
for row in ws_pat.iter_rows(min_row=2, values_only=True):
    if row[col["Study ID"]] == case_id:
        clip_start_ts = row[col[f"Point {clip_label} start"]]
        break

start_secs = ts_to_secs(str(clip_start_ts)) if clip_start_ts else 0.0
print(f"Clip {clip_label} window starts at {secs_to_ts(start_secs)}")

# Load force timestamps for this clip from points ABCD Excel
force_xlsx = BASE / "Videos" / case_id / "force_data" / f"{case_id}_points_ABCD_force.xlsx"
force_timestamps = []

if force_xlsx.exists():
    wb_force = openpyxl.load_workbook(force_xlsx)
    sheet_name = f"Point {clip_label}"
    if sheet_name in wb_force.sheetnames:
        ws_force = wb_force[sheet_name]
        force_headers = [cell.value for cell in ws_force[1]]
        in_window_col = force_headers.index("In Window")
        for row in ws_force.iter_rows(min_row=2, values_only=True):
            abs_sec = row[1]  # "Seconds" column
            in_window = row[in_window_col]
            if abs_sec is not None and in_window == "YES":
                force_timestamps.append(float(abs_sec))
        print(f"Found {len(force_timestamps)} force timestamps for clip {clip_label}")
    else:
        print(f"WARNING: Sheet '{sheet_name}' not found in {force_xlsx.name}")

if not force_timestamps:
    print("WARNING: No force timestamps found — falling back to 1-second intervals from clip start")
    force_timestamps = [start_secs + i for i in range(int(np.ceil(duration)) + 1)]

# Extract frames at force timestamps
for abs_sec in sorted(force_timestamps):
    if use_full_video:
        frame_num = int(abs_sec * fps)  # direct seek — no clip offset error
    else:
        frame_num = int((abs_sec - start_secs) * fps)
    if frame_num < 0 or frame_num >= total:
        continue
    cap.set(cv2.CAP_PROP_POS_FRAMES, frame_num)
    ret, frame = cap.read()
    if ret:
        frames[abs_sec] = (cv2.cvtColor(frame, cv2.COLOR_BGR2RGB), abs_sec)
cap.release()

print(f"Extracted {len(frames)} frames matching force data timestamps")

# --- Interactive marking ---
results = {}  # abs_sec -> {left, right, distance_px, distance_mm, timestamp}
frame_keys = sorted(frames.keys())
current_idx = [0]
frame_clicks = [[]]  # clicks for current frame

fig = plt.figure(figsize=(13, 10))
ax  = fig.add_axes([0, 0, 1, 1])
ax.axis("off")

info_text = ax.text(10, 10, "", fontsize=11, color="white", va="top",
                    bbox=dict(facecolor="black", alpha=0.7))

img_display = [None]
instruction_text = [None]  # persistent reference so onclick can update it

def show_frame(idx):
    ax.cla()
    ax.axis("off")
    abs_sec = frame_keys[idx]
    frame_img, _ = frames[abs_sec]
    ax.imshow(frame_img)
    img_display[0] = frame_img

    n_done = len(results)
    n_total = len(frame_keys)
    ax.set_title(
        f"Clip {clip_label}  |  Frame {idx+1}/{n_total}  |  t={secs_to_ts(abs_sec)}  |  "
        f"Done: {n_done}/{n_total}  |  SPACE=skip  ←/→=navigate  Q=quit",
        fontsize=12, color="white",
        bbox=dict(boxstyle="round", facecolor="#333333", alpha=0.85))

    if abs_sec in results:
        r = results[abs_sec]
        pl, pr = r["left_crus_px"], r["right_crus_px"]
        ax.plot(pl[0], pl[1], "+", color="#00FF88", markersize=14, markeredgewidth=2)
        ax.plot(pr[0], pr[1], "+", color="#FF6600", markersize=14, markeredgewidth=2)
        ax.plot([pl[0], pr[0]], [pl[1], pr[1]], "--", color="white", linewidth=1.5, alpha=0.8)
        ax.text((pl[0]+pr[0])//2, (pl[1]+pr[1])//2 - 15,
                f"{r['distance_mm']:.2f} mm",
                color="white", fontsize=9, ha="center",
                bbox=dict(facecolor="black", alpha=0.6))
        instruction_text[0] = ax.text(10, frame_img.shape[0] - 10,
                "Already marked — click to re-mark or press SPACE/→ to continue",
                fontsize=10, color="yellow", va="bottom",
                bbox=dict(facecolor="black", alpha=0.6))
    else:
        n_clicks = len(frame_clicks[0])
        msg = "Click 1/2: LEFT crus suture entry point" if n_clicks == 0 else "Click 2/2: RIGHT crus suture entry point"
        instruction_text[0] = ax.text(10, frame_img.shape[0] - 10, msg, fontsize=11,
                color="white", va="bottom", bbox=dict(facecolor="black", alpha=0.6))

    fig.canvas.draw()


def onclick(event):
    if event.xdata is None or event.ydata is None:
        return
    idx = current_idx[0]
    abs_sec = frame_keys[idx]
    x, y = int(event.xdata), int(event.ydata)

    # If already marked, clear and re-mark
    if abs_sec in results and len(frame_clicks[0]) == 0:
        del results[abs_sec]
        show_frame(idx)
        return

    frame_clicks[0].append((x, y))
    color = "#00FF88" if len(frame_clicks[0]) == 1 else "#FF6600"
    ax.plot(x, y, "+", color=color, markersize=14, markeredgewidth=2)

    # Update bottom instruction text after first click
    if len(frame_clicks[0]) == 1 and instruction_text[0] is not None:
        instruction_text[0].set_text("Click 2/2: RIGHT crus suture entry point")

    fig.canvas.draw()

    if len(frame_clicks[0]) == 2:
        p_left, p_right = frame_clicks[0]
        px_dist = np.sqrt((p_right[0]-p_left[0])**2 + (p_right[1]-p_left[1])**2)
        mm_dist = px_dist / px_per_mm

        ax.plot([p_left[0], p_right[0]], [p_left[1], p_right[1]],
                "--", color="white", linewidth=1.5, alpha=0.8)
        mid_x = (p_left[0] + p_right[0]) // 2
        mid_y = (p_left[1] + p_right[1]) // 2 - 15
        ax.annotate(f"{mm_dist:.2f} mm", (mid_x, mid_y), color="white", fontsize=10,
                    ha="center", bbox=dict(facecolor="black", alpha=0.7))
        ax.set_title(
            f"Clip {clip_label}  |  t={secs_to_ts(abs_sec)}  |  {mm_dist:.2f} mm  — press SPACE to confirm, click to redo",
            fontsize=12, color="white",
            bbox=dict(boxstyle="round", facecolor="#1a5276", alpha=0.85))
        fig.canvas.draw()

        results[abs_sec] = {
            "timestamp":     secs_to_ts(abs_sec),
            "abs_seconds":   round(abs_sec, 3),
            "left_crus_px":  list(p_left),
            "right_crus_px": list(p_right),
            "distance_px":   round(px_dist, 2),
            "distance_mm":   round(mm_dist, 4),
        }
        print(f"  t={secs_to_ts(abs_sec)}  {px_dist:.0f}px = {mm_dist:.2f}mm  [press SPACE to confirm]")

        frame_clicks[0] = []


def onkey(event):
    idx = current_idx[0]
    abs_sec = frame_keys[idx]

    if event.key == " " or event.key == "right":
        # Confirm marked frame and advance, or skip unmarked frame
        frame_clicks[0] = []
        if abs_sec in results:
            print(f"  Confirmed t={secs_to_ts(abs_sec)}")
        else:
            print(f"  Skipped t={secs_to_ts(abs_sec)}")
        if idx + 1 < len(frame_keys):
            current_idx[0] += 1
            show_frame(current_idx[0])
        else:
            finish()
    elif event.key == "left":
        frame_clicks[0] = []
        if idx > 0:
            current_idx[0] -= 1
            show_frame(current_idx[0])
    elif event.key == "q":
        finish()


finished = [False]

def finish():
    if finished[0]:
        return
    finished[0] = True
    print(f"\nMarked {len(results)}/{len(frame_keys)} frames")
    save_results()
    # Redraw last frame cleanly with all markers
    show_frame(current_idx[0])
    ax.set_title(f"Done — {len(results)} frames marked. Close window.",
                 fontsize=13, color="white",
                 bbox=dict(boxstyle="round", facecolor="#1e8449", alpha=0.8))
    fig.canvas.draw()


def save_results():
    out_json = BASE / "Videos" / case_id / "force_data" / f"{case_id}_suture_entry.json"
    out_xlsx = BASE / "Videos" / case_id / "force_data" / f"{case_id}_suture_entry.xlsx"

    # Load existing data
    data = {}
    if out_json.exists():
        with open(out_json) as f:
            data = json.load(f)

    data[clip_label] = {
        "px_per_mm": px_per_mm,
        "frames": results,
    }

    with open(out_json, "w") as f:
        json.dump(data, f, indent=2)
    print(f"Saved JSON: {out_json}")

    # Excel
    if out_xlsx.exists():
        wb = openpyxl.load_workbook(out_xlsx)
        # Remove existing sheet for this clip
        if clip_label in wb.sheetnames:
            del wb[clip_label]
    else:
        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

    ws = wb.create_sheet(title=clip_label)
    ws.append(["Timestamp", "Abs seconds", "Left crus (px)", "Right crus (px)",
               "Distance (px)", "Distance (mm)", "px/mm"])

    for abs_sec in sorted(results.keys()):
        r = results[abs_sec]
        ws.append([
            r["timestamp"],
            r["abs_seconds"],
            str(r["left_crus_px"]),
            str(r["right_crus_px"]),
            r["distance_px"],
            r["distance_mm"],
            px_per_mm,
        ])

    wb.save(out_xlsx)
    print(f"Saved Excel: {out_xlsx}")


fig.canvas.mpl_connect("button_press_event", onclick)
fig.canvas.mpl_connect("key_press_event", onkey)
plt.get_current_fig_manager().full_screen_toggle()
show_frame(0)
plt.show()
