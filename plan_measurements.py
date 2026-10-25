"""
plan_measurements.py — Summarise what can be measured and set the stride for each point.

Run after extract_point_forces.py. Each selected force second (window and buffer rows) is measured on the
frames around its centre frame, the frame whose stored time is nearest the sample's time. This script prints,
per point and in total:
    - how many force values were selected (window and buffer)
    - the maximum number of frames available (every frame inside the limit: stride 1)
    - how many frames (and clicks, two per frame) other strides would need
and then asks for the stride of each point: 1 measures every frame, 10 measures every tenth frame, and so
on. The choice is saved in a measurement plan that mark_suture_entry.py reads.

Reads:
    Videos/<case_id>/force_data/<case_id>_frame_times.csv            (from index_video.py)
    Videos/<case_id>/force_data/<case_id>_points_ABCD_force.xlsx     (from extract_point_forces.py)
Writes:
    Videos/<case_id>/force_data/<case_id>_measurement_plan.json

No frame serves two force seconds. Force samples are one second apart (30 frames at 30 fps), and a frame
exactly half that far from the centre frame is halfway to the neighbouring sample. Every offset therefore
stays strictly less than half the spacing: at most 14 frames either side at 30 fps (about 0.47 s), so one
distance measurement feeds one point of the force-strain curve. The offsets are 0, +/-s, +/-2s, ... within
that limit: stride 10 gives -10, 0, +10; stride 5 gives -10 to +10 by 5; stride 1 gives -14 to +14. The
default half-window is the largest allowed; --half-window-s may lower it but a value that would allow 15
frames (0.5 s at 30 fps) or more is refused. All of this assumes that a force sample's time is the centre of
its second, not its start (see the README, "Timing convention").
Marks are stored by frame, so the plan can be re-run later with a smaller stride to add frames; frames
already marked are kept.

Usage:
    python3 plan_measurements.py <case_id> [--stride A=10 B=5 ...] [--half-window-s S]
            [--window-center-shift-s S] [--yes]
    --stride accepts one number for every point (--stride 10) or per-point values (--stride A=10 B=1).
    Without --stride the script asks for each point; --yes accepts the default (10) without asking.
"""

import argparse
import json
import sys
from pathlib import Path

import framelib as fl

BASE = Path(__file__).resolve().parent
POINTS = ["A", "B", "C", "D"]
DEFAULT_STRIDE = 10
COMPARE_STRIDES = (1, 5, 10)
SAMPLING_KEYS = ("half_window_frames", "time_offset_s", "window_center_shift_s")


def read_selection_info(workbook_path):
    """Buffer and time offset recorded by extract_point_forces.py (None, 0.0 if the workbook has no record)."""
    import openpyxl
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "Selection" not in wb.sheetnames:
            return {"buffer_s": None, "time_offset_s": 0.0}
        rows = wb["Selection"].iter_rows(values_only=True)
        headers = list(next(rows))
        first = next((r for r in rows if r[0] in POINTS), None)
    finally:
        wb.close()
    if first is None:
        return {"buffer_s": None, "time_offset_s": 0.0}
    return {"buffer_s": float(first[headers.index("Buffer (s)")]),
            "time_offset_s": float(first[headers.index("Time offset (s)")])}


def point_summary(index, rows, params):
    """Counts for one point: force values by role, maximum frames, frames at the comparison strides."""
    seconds = [x for x, _ in rows]
    base = {k: params[k] for k in SAMPLING_KEYS}
    every_frame = fl.plan_point(index, seconds, stride=1, **base)
    out = {
        "force_values": len(rows),
        "window_values": sum(1 for _, r in rows if r == "window"),
        "buffer_values": sum(1 for _, r in rows if r != "window"),
        "max_frames": len(every_frame.frames),
        "shared_frames": len(every_frame.shared_frames()),
    }
    out["frames_at"] = {s: len(fl.plan_point(index, seconds, stride=s, **base).frames) for s in COMPARE_STRIDES}
    return out


def parse_strides(tokens, labels):
    """--stride tokens -> {label: stride}. One bare number applies to every point; 'A=10' sets one point."""
    result = {}
    for tok in tokens or []:
        if "=" in tok:
            label, val = tok.split("=", 1)
            label = label.strip().upper()
            if label not in labels:
                raise ValueError(f"unknown point {label!r} in --stride (points: {', '.join(labels)})")
            result[label] = int(val)
        else:
            for label in labels:
                result.setdefault(label, int(tok))
    for label, s in result.items():
        if s < 1:
            raise ValueError(f"stride for point {label} must be 1 or more")
    return result


def ask_strides(labels, given, summaries, input_fn=input, interactive=True, assume_yes=False):
    """Stride per point: from --stride, else asked at the terminal, else the default."""
    strides = dict(given)
    for label in labels:
        if label in strides:
            continue
        if not interactive or assume_yes:
            strides[label] = DEFAULT_STRIDE
            continue
        s = summaries[label]
        while True:
            reply = input_fn(f"Stride for point {label} (1 = every frame, default {DEFAULT_STRIDE}): ").strip()
            if reply == "":
                strides[label] = DEFAULT_STRIDE
                break
            try:
                value = int(reply)
                if value >= 1:
                    strides[label] = value
                    break
            except ValueError:
                pass
            print("  please enter a whole number of 1 or more")
    return strides


def main(argv=None, input_fn=input, interactive=None):
    ap = argparse.ArgumentParser(description="Summarise measurable frames and set the stride for each point.")
    ap.add_argument("case_id")
    ap.add_argument("--stride", nargs="+", default=None,
                    help="one number for all points, or per point, e.g. A=10 B=5")
    ap.add_argument("--half-window-s", type=float, default=None, dest="half_window_s",
                    help="half-window around each force second, in seconds. Must stay below half the spacing of the "
                         "force samples (under 0.5 s: at most 14 frames at 30 fps) so that no frame serves two force "
                         "seconds; a larger value is refused [the largest allowed]")
    ap.add_argument("--window-center-shift-s", type=float, default=0.0, dest="window_center_shift_s",
                    help="shift the window centre. The pipeline assumes a force sample at second x is the centre of "
                         "that second; if the console means the second that starts at x, the frames sit half a "
                         "second early and +0.5 corrects it. Not verified: see the README, Timing convention [0]")
    ap.add_argument("--yes", action="store_true", help="accept the default stride without asking")
    ap.add_argument("--version", action="version", version=f"%(prog)s {fl.PIPELINE_VERSION}")
    args = ap.parse_args(argv)
    interactive = sys.stdin.isatty() if interactive is None else interactive

    case_id = args.case_id
    force_dir = BASE / "Videos" / case_id / "force_data"
    workbook = force_dir / f"{case_id}_points_ABCD_force.xlsx"
    if not workbook.exists():
        sys.exit(f"ERROR: {workbook.name} not found - run extract_point_forces.py first")
    index = fl.load_index(force_dir, case_id)

    import openpyxl
    wb_sheets = openpyxl.load_workbook(workbook, read_only=True)
    sheets = list(wb_sheets.sheetnames)
    wb_sheets.close()
    labels = [p for p in POINTS if f"Point {p}" in sheets]
    if not labels:
        sys.exit(f"ERROR: no point sheets in {workbook.name}")
    info = read_selection_info(workbook)

    if args.half_window_s is None:
        half_frames = fl.max_half_window_frames(index)
    else:
        half_frames = fl.half_window_frames_from_seconds(index, args.half_window_s)
    half_window_s = round(half_frames * index.mean_gap_s, 3) if args.half_window_s is None else args.half_window_s
    try:
        fl.check_half_window(index, half_frames)
    except ValueError as e:
        sys.exit(f"ERROR: {e}")
    params = {"half_window_frames": half_frames,
              "time_offset_s": info["time_offset_s"], "window_center_shift_s": args.window_center_shift_s}

    rows_by_point, summaries, legacy_any = {}, {}, False
    for label in labels:
        rows, legacy = fl.selected_force_rows(workbook, label)
        legacy_any |= legacy
        rows_by_point[label] = rows
        summaries[label] = point_summary(index, rows, params)

    print(f"\nMeasurement plan for case {case_id}   (pipeline v{fl.PIPELINE_VERSION})")
    print(f"Frame table: {index.n} frames, mean {1.0 / index.mean_gap_s:.6f} fps.  "
          f"Window around each force second: +/-{half_frames} frames (+/-{half_window_s:g} s); the limit is "
          f"{fl.max_half_window_frames(index)} frames, so no frame serves two force seconds.")
    print(f"Buffer: {'not recorded' if info['buffer_s'] is None else format(info['buffer_s'], 'g') + ' s'}; "
          f"time offset (video minus force): {info['time_offset_s']:g} s.")
    if legacy_any:
        print("NOTE: the force workbook has no Role column (made by an older extraction); only its 'In Window' rows "
              "are counted. Re-run extract_point_forces.py to include buffer rows.")
    head = (f"{'Point':<6}{'Force values':>13}{'window':>8}{'buffer':>8}{'Max frames':>12}   "
            + "".join(f"{'stride ' + str(s):>12}" for s in COMPARE_STRIDES))
    print("\n" + head)
    print("-" * len(head))
    totals = {"force": 0, "window": 0, "buffer": 0, "max": 0, **{s: 0 for s in COMPARE_STRIDES}}
    for label in labels:
        s = summaries[label]
        print(f"{label:<6}{s['force_values']:>13}{s['window_values']:>8}{s['buffer_values']:>8}{s['max_frames']:>12}   "
              + "".join(f"{s['frames_at'][k]:>12}" for k in COMPARE_STRIDES))
        totals["force"] += s["force_values"]
        totals["window"] += s["window_values"]
        totals["buffer"] += s["buffer_values"]
        totals["max"] += s["max_frames"]
        for k in COMPARE_STRIDES:
            totals[k] += s["frames_at"][k]
    print("-" * len(head))
    print(f"{'Total':<6}{totals['force']:>13}{totals['window']:>8}{totals['buffer']:>8}{totals['max']:>12}   "
          + "".join(f"{totals[k]:>12}" for k in COMPARE_STRIDES))
    shared = sum(summaries[label]["shared_frames"] for label in labels)
    print(f"Frames used for two force seconds: {shared}.  Each frame takes two clicks.\n")
    if shared:
        sys.exit("ERROR: some frames would serve two force seconds (force samples closer together than the "
                 "half-window allows, for example a video with dropped frames). Lower --half-window-s.")

    try:
        given = parse_strides(args.stride, labels)
    except ValueError as e:
        sys.exit(f"ERROR: {e}")
    if not interactive and not args.yes and len(given) < len(labels):
        print(f"(not an interactive terminal: using the default stride {DEFAULT_STRIDE} for points without --stride)")
    strides = ask_strides(labels, given, summaries, input_fn, interactive, args.yes)

    old = fl.load_plan(force_dir, case_id, check_version=False)
    points = {}
    print("\nChosen:")
    for label in labels:
        s = strides[label]
        frames = len(fl.plan_point(index, [x for x, _ in rows_by_point[label]], stride=s,
                                   **{k: params[k] for k in SAMPLING_KEYS}).frames)
        if s > half_frames:
            print(f"  note: stride {s} is wider than the half-window ({half_frames} frames); "
                  f"point {label} gets one frame per force second")
        points[label] = {**{k: summaries[label][k] for k in ("force_values", "window_values", "buffer_values", "max_frames")},
                         "stride": s, "planned_frames": frames}
        print(f"  Point {label}: stride {s} -> {frames} frames ({2 * frames} clicks)")
    total_frames = sum(p["planned_frames"] for p in points.values())
    print(f"  Total: {total_frames} frames ({2 * total_frames} clicks)")

    plan = {
        "schema_version": fl.PLAN_SCHEMA_VERSION, "pipeline_version": fl.PIPELINE_VERSION, "case_id": case_id,
        "created": fl.utc_now(),
        "frame_table": {"video_name": index.source_name, "frames": index.n, "mean_fps": round(1.0 / index.mean_gap_s, 6)},
        "parameters": {"half_window_s": half_window_s, "half_window_frames": half_frames,
                       "window_center_shift_s": args.window_center_shift_s,
                       "time_offset_s": info["time_offset_s"], "buffer_s": info["buffer_s"]},
        "points": points,
    }
    out = fl.plan_path(force_dir, case_id)
    if old:
        print(f"\nReplacing the plan made {old.get('created', '?')} by version {old.get('pipeline_version', '?')}.")
    tmp = out.with_name(out.name + ".tmp")
    with open(tmp, "w") as f:
        json.dump(plan, f, indent=1)
    tmp.replace(out)
    print(f"Saved: {out}")
    print("Next:\n"
          f"  1. the ruler frame, once:  python3 export_calibration_frame.py <video> {case_id} ruler --time HH:MM:SS\n"
          f"                             python3 calibrate_reference.py <the PNG it saved> {case_id}\n"
          "  2. each point A to D:      python3 pick_calibration_frame.py <video> " + case_id + " <point>   (scrub, Enter on the frame to use)\n"
          f"                             python3 calibrate_reference.py <the PNG it saved> {case_id} <point>\n"
          f"  3. then each point:        python3 mark_suture_entry.py <video> {case_id} <point>")


if __name__ == "__main__":
    main()
