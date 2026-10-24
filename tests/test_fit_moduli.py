"""
End-to-end test on a synthetic case: per-frame marks -> merge_distance_force.py ->
fit_moduli.py. The gap closes linearly and force rises linearly with strain, so the
fitted linear modulus can be checked against a plain regression on the per-second means,
for both row sets (--rows all and --rows window).

Each point has 2 buffer-before rows, 8 window rows and 2 buffer-after rows.

Run from the repository root:
    python3 -m unittest discover -s tests -v
"""

import sys
import tempfile
import unittest
from pathlib import Path

import numpy as np
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import framelib as fl  # noqa: E402
import fit_moduli as fm  # noqa: E402
import merge_distance_force as mg  # noqa: E402
from test_framelib import synthetic_index  # noqa: E402

CASE = "case1"
PXMM = 10.0
STARTS = {"A": 8378, "B": 8413, "C": 8453, "D": 8490}
N_SEC = 8
BUF = 2


def gap(t, start):
    return 40.0 - 4.0 * (t - start)        # mm, closing at 4 mm/s


def rows_for(start, leftover=False):
    """(force second, role, USM1) for one point. leftover: 1.7 N still on the first buffer row, from earlier activity."""
    out = []
    for j in range(-BUF, N_SEC + BUF):
        role = "buffer-before" if j < 0 else ("window" if j < N_SEC else "buffer-after")
        f1 = 0.0 if j < 2 else 0.8 * (j - 1)          # zero force until the 2nd window second, then rising
        if leftover and j == -BUF:
            f1 = 1.7
        out.append((start + j, role, f1))
    return out


def build_case(root, idx, params, leftover=False, arms=("USM1", "USM3")):
    fdir = Path(root) / "Videos" / CASE / "force_data"
    fdir.mkdir(parents=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    store = fl.MarksStore(fdir / f"{CASE}_frame_marks.json", CASE)
    store.set_video(idx)
    for label, start in STARTS.items():
        ws = wb.create_sheet(f"Point {label}")
        ws.append(["Timestamp", "Seconds", *arms, "In Window", "Role"])
        rows = rows_for(start, leftover)
        for x, role, f1 in rows:       # the first arm carries f1, every other arm half of it
            ws.append([fl.fmt_time(x), float(x), *[f1 if i == 0 else 0.5 * f1 for i in range(len(arms))],
                       "YES" if role == "window" else "", role])
        store.note_params(label, params)
        plan = fl.plan_point(idx, [x for x, _, _ in rows], **params)
        for k in plan.frames:
            d = gap(idx.pts_s[k], start)
            store.put(label, k, {"pts_s": float(idx.pts_s[k]), "left_crus_px": [0, 0], "right_crus_px": [d * PXMM, 0],
                                 "distance_px": d * PXMM, "distance_mm": d, "px_per_mm": PXMM,
                                 "skipped": False, "source": "manual"})
    wb.save(fdir / f"{CASE}_points_ABCD_force.xlsx")
    idx.save(fl.frame_table_path(fdir, CASE))
    store.save()
    return fdir


def merged_rows(fdir, label="A"):
    ws = openpyxl.load_workbook(fdir / f"{CASE}_points_ABCD_force_merged.xlsx")[f"Point {label}"]
    head = [c.value for c in ws[1]]
    return [dict(zip(head, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]


def expected_slope(fdir, rows_mode):
    data = [r for r in merged_rows(fdir) if r["Distance (mm)"] is not None
            and (rows_mode == "all" or r["Role"] == "window")]
    d = np.array([r["Distance (mm)"] for r in data], float)
    f = np.array([r["Total Force (N)"] for r in data], float)
    eps = (d.max() - d) / d.max()
    return np.polyfit(eps, f, 1)[0], len(data)


class TestConventionB(unittest.TestCase):
    """The zero-strain reference for Convention B, including force traces that return to zero."""

    dist = np.array([40.0, 38, 34, 30, 26, 22, 18, 14, 12, 12])
    # zero before loading, loading, then the force falls back to zero when the instrument is released
    force = np.array([0.0, 0, 0, 1.0, 2.0, 3.0, 4.0, 3.0, 0.0, 0.0])

    def test_first_rise_takes_the_last_zero_before_loading(self):
        self.assertEqual(fm.onset_index(self.dist, self.force, "first-rise"), 2)
        cb = fm.convention_b(self.dist, self.force, 40.0, "first-rise")
        self.assertEqual((cb["oi"], cb["d0"], cb["L0"]), (2, 6.0, 34.0))
        self.assertEqual(int(cb["keep"].sum()), 8)                    # rows 2 to 9

    def test_version_1_rule_jumps_to_the_end_when_force_returns_to_zero(self):
        # the failure that motivated the change: the smallest-gap zero-force row is the post-release row
        self.assertEqual(fm.onset_index(self.dist, self.force, "smallest-gap"), 8)
        cb = fm.convention_b(self.dist, self.force, 40.0, "smallest-gap")
        self.assertEqual(int(cb["keep"].sum()), 2)

    def test_both_rules_agree_when_force_is_zero_only_before_loading(self):
        force = np.array([0.0, 0, 0, 1.0, 2.0, 3.0, 4.0, 5.0, 6.0, 7.0])
        self.assertEqual(fm.onset_index(self.dist, force, "first-rise"),
                         fm.onset_index(self.dist, force, "smallest-gap"))
        a = fm.convention_b(self.dist, force, 40.0, "first-rise")
        b = fm.convention_b(self.dist, force, 40.0, "smallest-gap")
        self.assertEqual(a["d0"], b["d0"])
        self.assertTrue(np.array_equal(a["keep"], b["keep"]))

    # roles for a 10-row point: the first 4 rows are buffer, the window starts at row 4
    roles = np.array(["buffer-before"] * 4 + ["window"] * 6)

    def test_window_rise_ignores_leftover_force_in_the_buffer(self):
        # point D of the real test case: 1.7 N left over on the first row, zeros after it, the rise one row after the window start
        force = np.array([1.7, 0, 0, 0, 0, 1.0, 2.0, 3.0, 4.0, 5.0])
        self.assertIsNone(fm.onset_index(self.dist, force, "first-rise"))               # stops at the leftover row
        self.assertEqual(fm.onset_index(self.dist, force, "window-rise", self.roles), 4)
        cb = fm.convention_b(self.dist, force, 40.0, "window-rise", self.roles)
        self.assertEqual((cb["oi"], int(cb["keep"].sum())), (4, 6))

    def test_window_rise_may_put_the_reference_in_the_buffer(self):
        # force already rising in the buffer when the window starts: the last zero row before the
        # first nonzero row at the window start is the buffer row before the buffer's own rise
        force = np.array([0.0, 0, 0.5, 0.6, 0.7, 1.0, 2.0, 3.0, 4.0, 5.0])
        self.assertEqual(fm.onset_index(self.dist, force, "window-rise", self.roles), 1)
        self.assertEqual(fm.onset_index(self.dist, force, "first-rise"), 1)

    def test_window_rise_uses_the_rise_at_the_noted_start_not_an_earlier_blip(self):
        # a 0.8 N blip in the buffer, then zeros, then the real rise at the window start (point C of the test case)
        force = np.array([0.0, 0.8, 0, 0, 0, 0, 1.0, 2.0, 3.0, 4.0])
        self.assertEqual(fm.onset_index(self.dist, force, "first-rise"), 0)             # before the blip
        self.assertEqual(fm.onset_index(self.dist, force, "window-rise", self.roles), 5)

    def test_window_rise_undefined_cases(self):
        self.assertIsNone(fm.onset_index(self.dist, np.ones(10), "window-rise", self.roles))   # nonzero from row 0 on
        force = np.array([1.0, 1, 1, 1, 1, 1.0, 2.0, 3.0, 4.0, 5.0])                           # nonzero from the first row to the end
        self.assertIsNone(fm.onset_index(self.dist, force, "window-rise", self.roles))
        zeros_from_start = np.array([0.5, 0.5, 0, 0, 0, 0, 0, 0, 0, 0])                        # no rise from the window start on
        self.assertIsNone(fm.onset_index(self.dist, zeros_from_start, "window-rise", self.roles))
        self.assertIsNone(fm.onset_index(self.dist, np.zeros(10), "window-rise", self.roles))  # never nonzero
        no_window = np.array(["buffer-before"] * 10)
        self.assertIsNone(fm.onset_index(self.dist, self.force, "window-rise", no_window))      # no window row at all

    def test_window_rise_equals_first_rise_when_every_row_is_a_window_row(self):
        # --rows window: the rows start at the noted start, so the two rules agree
        for force in (self.force, np.array([0.0, 0, 0, 1.0, 2.0, 3.0, 4.0, 5.0, 6.0, 7.0]), np.ones(10)):
            expected = fm.onset_index(self.dist, force, "first-rise")
            self.assertEqual(fm.onset_index(self.dist, force, "window-rise"), expected)
            self.assertEqual(fm.onset_index(self.dist, force, "window-rise", np.array(["window"] * 10)), expected)

    def test_window_rise_is_the_default_rule(self):
        self.assertEqual(fm.ONSET_RULES[0], "window-rise")
        self.assertEqual(fm.parse_args(["case1"]).onset, "window-rise")
        self.assertEqual(fm.parse_args(["case1", "--onset", "smallest-gap"]).onset, "smallest-gap")

    def test_undefined_cases(self):
        self.assertIsNone(fm.onset_index(self.dist, np.ones(10), "first-rise"))          # nonzero from the first row
        self.assertIsNone(fm.onset_index(self.dist, np.zeros(10), "first-rise"))         # never nonzero
        self.assertIsNone(fm.onset_index(self.dist, np.ones(10), "smallest-gap"))        # no zero at all
        self.assertIsNone(fm.convention_b(self.dist, np.ones(10), 40.0))


class TestFitChain(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.idx = synthetic_index()

    def run_chain(self, root, *fit_args):
        mg.BASE = Path(root)
        mg.main([CASE])
        fm.BASE, fm.case_id = Path(root), CASE
        fm.main(list(fit_args))

    def table_row(self, fdir, rows_mode, point="A", channel="Total"):
        table = openpyxl.load_workbook(fdir / f"{CASE}_moduli_table_{rows_mode}.xlsx").active
        head = [c.value for c in table[1]]
        for r in table.iter_rows(min_row=2, values_only=True):
            if r[0] == point and r[1] == channel:
                return dict(zip(head, r))
        raise AssertionError("row not found")

    def test_all_rows_is_the_default_and_matches_a_plain_regression(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir = build_case(root, self.idx, params)
            self.run_chain(root)
            plots = Path(root) / "Videos" / CASE / "plots"
            for p in "ABCD":
                self.assertGreater((plots / f"{CASE}_point{p}_model_fits_all.png").stat().st_size, 10_000)
            row = self.table_row(fdir, "all")
            slope, n = expected_slope(fdir, "all")
            self.assertEqual((row["n"], row["n window"], row["n buffer"], row["Rows used"]), (n, N_SEC, 2 * BUF, "all"))
            self.assertAlmostEqual(row["A: Linear E (N)"], round(slope, 3), places=3)

    def test_window_rows_option_leaves_buffer_out_and_names_files_apart(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir = build_case(root, self.idx, params)
            self.run_chain(root)
            fm.main(["--rows", "window"])
            plots = Path(root) / "Videos" / CASE / "plots"
            self.assertTrue((plots / f"{CASE}_pointA_model_fits_window.png").exists())
            self.assertTrue((plots / f"{CASE}_pointA_model_fits_all.png").exists())          # both coexist
            row = self.table_row(fdir, "window")
            slope, n = expected_slope(fdir, "window")
            self.assertEqual((row["n"], row["n window"], row["n buffer"], row["Rows used"]), (N_SEC, N_SEC, 0, "window"))
            self.assertAlmostEqual(row["A: Linear E (N)"], round(slope, 3), places=3)
            self.assertNotEqual(self.table_row(fdir, "all")["n"], row["n"])

    def test_row_set_changes_the_strain_reference(self):
        # the widest gap lies in the buffer-before rows, so 'all' and 'window' use different references
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir = build_case(root, self.idx, params)
            self.run_chain(root)
            fm.main(["--rows", "window"])
            self.assertNotAlmostEqual(self.table_row(fdir, "all")["A: Linear E (N)"],
                                      self.table_row(fdir, "window")["A: Linear E (N)"], places=2)

    def test_table_footer_names_rows_and_version(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir = build_case(root, self.idx, params)
            self.run_chain(root, "--rows", "window")
            table = openpyxl.load_workbook(fdir / f"{CASE}_moduli_table_window.xlsx").active
            text = " ".join(str(c.value) for row in table.iter_rows() for c in row if c.value is not None)
            self.assertIn("window rows only", text)
            self.assertIn(fl.PIPELINE_VERSION, text)
            self.assertIn("Convention B reference rule: window-rise", text)

    def test_channels_follow_the_arms_in_the_workbook(self):
        params = fl.resolve_params(stride=10)
        for arms in (("USM1", "USM2"), ("USM2", "USM4"), ("USM1", "USM2", "USM3", "USM4")):
            with self.subTest(arms=arms), tempfile.TemporaryDirectory() as root:
                fdir = build_case(root, self.idx, params, arms=arms)
                self.run_chain(root)
                table = openpyxl.load_workbook(fdir / f"{CASE}_moduli_table_all.xlsx").active
                channels = {r[1] for r in table.iter_rows(min_row=2, values_only=True) if r[0] in list("ABCD")}
                self.assertEqual(channels, set(arms) | {"Total"})
                first, second = (self.table_row(fdir, "all", "A", a)["A: Linear E (N)"] for a in arms[:2])
                self.assertAlmostEqual(second, 0.5 * first, places=2)          # each arm reads its own column
                total = self.table_row(fdir, "all", "A", "Total")["A: Linear E (N)"]
                self.assertAlmostEqual(total, first * (1 + 0.5 * (len(arms) - 1)), places=2)
                for p in "ABCD":
                    self.assertGreater((Path(root) / "Videos" / CASE / "plots"
                                        / f"{CASE}_point{p}_model_fits_all.png").stat().st_size, 10_000)

    def test_default_rule_survives_a_leftover_nonzero_buffer_row(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir = build_case(root, self.idx, params, leftover=True)
            self.run_chain(root)                                            # default: window-rise
            row = self.table_row(fdir, "all", "A", "USM1")
            self.assertIsNotNone(row["B: Linear E (N)"])
            self.assertIn("onset@", row["B note"])
            self.run_chain(root, "--onset", "first-rise")                   # the earlier rule stops at the leftover row
            row = self.table_row(fdir, "all", "A", "USM1")
            self.assertIsNone(row["B: Linear E (N)"])
            self.assertIn("B undefined", row["B note"])
            table = openpyxl.load_workbook(fdir / f"{CASE}_moduli_table_all.xlsx").active
            text = " ".join(str(c.value) for r in table.iter_rows() for c in r if c.value is not None)
            self.assertIn("Convention B reference rule: first-rise", text)

    def test_rows_window_gives_the_same_reference_under_both_rules(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir = build_case(root, self.idx, params, leftover=True)
            self.run_chain(root, "--rows", "window")
            a = self.table_row(fdir, "window", "A", "USM1")
            self.run_chain(root, "--rows", "window", "--onset", "first-rise")
            b = self.table_row(fdir, "window", "A", "USM1")
            self.assertEqual((a["B: Linear E (N)"], a["B note"]), (b["B: Linear E (N)"], b["B note"]))

    def test_plots_still_work_without_the_frame_table(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir = build_case(root, self.idx, params)
            mg.BASE = Path(root)
            mg.main([CASE])
            (fdir / f"{CASE}_frame_distances.csv").unlink()
            fm.BASE, fm.case_id = Path(root), CASE
            fm.main()
            self.assertTrue((Path(root) / "Videos" / CASE / "plots" / f"{CASE}_pointA_model_fits_all.png").exists())

    def test_fit_asks_for_merge_first(self):
        with tempfile.TemporaryDirectory() as root:
            fm.BASE, fm.case_id = Path(root), CASE
            with self.assertRaises(SystemExit):
                fm.main()


if __name__ == "__main__":
    unittest.main()
