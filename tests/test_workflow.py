"""
Tests for the stages that lead up to measurement: extract_point_forces.py (selection rule),
plan_measurements.py (counts, strides, plan file) and the plan/role handling of
mark_suture_entry.py (--list, no display needed).

Run from the repository root:
    python3 -m unittest discover -s tests -v
"""

import contextlib
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import extract_point_forces as ex  # noqa: E402
import framelib as fl  # noqa: E402
import mark_suture_entry as ms  # noqa: E402
import plan_measurements as pm  # noqa: E402
from test_framelib import synthetic_index  # noqa: E402

CASE = "case1"


class TestSelectRows(unittest.TestCase):
    def test_roles_and_inclusive_bounds(self):
        rows = ex.select_rows(range(90, 120), 100, 105, buffer_s=5)
        roles = dict(rows)
        self.assertEqual([x for x, r in rows if r == "buffer-before"], [95, 96, 97, 98, 99])
        self.assertEqual([x for x, r in rows if r == "window"], [100, 101, 102, 103, 104, 105])
        self.assertEqual([x for x, r in rows if r == "buffer-after"], [106, 107, 108, 109, 110])
        self.assertNotIn(94, roles)
        self.assertNotIn(111, roles)

    def test_zero_buffer_keeps_only_the_window(self):
        rows = ex.select_rows(range(90, 120), 100, 105, buffer_s=0)
        self.assertEqual({r for _, r in rows}, {"window"})
        self.assertEqual(len(rows), 6)

    def test_sparse_force_log_gives_only_existing_seconds(self):
        rows = ex.select_rows([96, 99, 103, 108, 130], 100, 105, buffer_s=5)
        self.assertEqual(rows, [(96, "buffer-before"), (99, "buffer-before"), (103, "window"), (108, "buffer-after")])

    def test_time_offset_converts_force_time_to_video_time(self):
        # video time = force time + offset; offset -2 means force second 102 is at video time 100
        rows = ex.select_rows(range(90, 120), 100, 105, buffer_s=0, time_offset_s=-2.0)
        self.assertEqual([x for x, _ in rows], [102, 103, 104, 105, 106, 107])

    def test_whole_second_times_and_fractions_parse(self):
        self.assertEqual(ex.ts_to_secs("02:19:38"), 8378.0)
        self.assertEqual(ex.ts_to_secs("02:19:38.250"), 8378.25)


class TestExtractMain(unittest.TestCase):
    """extract_point_forces.main() on a synthetic force log and tracking sheet (any case ID is accepted)."""

    def build(self, root, cases=("CASE001",)):
        root = Path(root)
        sheet = openpyxl.Workbook()
        ws = sheet.active
        ws.append(["Study ID"] + [f"Point {p} {side}" for p in "ABCD" for side in ("start", "end")])
        for c in cases:
            ws.append([c, "00:01:40", "00:01:45"] + [None] * 6)      # whole-second window 100-105; points B-D not measured
        ws.append([None, "00:00:10", "00:00:20"] + [None] * 6)        # blank study ID is skipped
        sheet.save(root / "sheet.xlsx")
        for c in cases:
            fd = root / "Videos" / c / "force_data"
            fd.mkdir(parents=True)
            usm1 = [{"x": x, "y": 1.0, "z": None} for x in range(90, 120)]
            usm3 = [{"x": x, "y": 0.5, "z": None} for x in range(90, 120) if x % 2 == 0]
            (fd / f"{c}_force_raw.json").write_text(json.dumps({"USM1": usm1, "USM2": [], "USM3": usm3, "USM4": []}))
        ex.BASE = root
        return root / "sheet.xlsx"

    def extract(self, root, sheet, *argv):
        return run_quiet(ex.main, [*argv, "--tracking-sheet", str(sheet)])

    def read(self, root, case="CASE001"):
        return openpyxl.load_workbook(Path(root) / "Videos" / case / "force_data" / f"{case}_points_ABCD_force.xlsx")

    def test_roles_buffer_fills_and_selection_sheet(self):
        with tempfile.TemporaryDirectory() as root:
            sheet = self.build(root)
            self.extract(root, sheet, "CASE001")
            wb = self.read(root)
            self.assertEqual(wb.sheetnames, ["Point A", "Selection"])                # no point B: not written
            ws = wb["Point A"]
            head = [c.value for c in ws[1]]
            self.assertEqual(head, ["Timestamp", "Seconds", "USM1", "USM3", "In Window", "Role"])
            rows = [dict(zip(head, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
            by_role = {}
            for r in rows:
                by_role.setdefault(r["Role"], []).append(r["Seconds"])
            self.assertEqual(by_role["buffer-before"], [95.0, 96, 97, 98, 99])
            self.assertEqual(by_role["window"], [100.0, 101, 102, 103, 104, 105])
            self.assertEqual(by_role["buffer-after"], [106.0, 107, 108, 109, 110])
            self.assertTrue(all((r["In Window"] == "YES") == (r["Role"] == "window") for r in rows))
            fills = {r["Role"]: ws.cell(row=2 + rows.index(r), column=1).fill.start_color.rgb for r in rows}
            self.assertNotEqual(fills["window"], fills["buffer-before"])
            self.assertEqual(fills["buffer-before"], fills["buffer-after"])
            sel = [[c.value for c in r] for r in wb["Selection"].iter_rows(min_row=1, max_row=2)]
            self.assertEqual(sel[1][:9], ["A", "00:01:40", "00:01:45", 5.0, 0.0, 6, 5, 5, "95 to 110"])
            self.assertEqual(rows[0]["USM3"], None)                                  # odd second: no USM3 sample

    def test_buffer_and_offset_options(self):
        with tempfile.TemporaryDirectory() as root:
            sheet = self.build(root)
            self.extract(root, sheet, "CASE001", "--buffer-s", "0", "--time-offset-s", "-2")
            ws = self.read(root)["Point A"]
            secs = [r[1].value for r in ws.iter_rows(min_row=2)]
            self.assertEqual(secs, [102.0, 103, 104, 105, 106, 107])                 # video 100-105 is force 102-107
            sel = [c.value for c in self.read(root)["Selection"][2]]
            self.assertEqual((sel[3], sel[4]), (0.0, -2.0))

    def test_case_filter_any_id_and_blank_rows(self):
        with tempfile.TemporaryDirectory() as root:
            sheet = self.build(root, cases=("CASE001", "SITE2-0042"))
            out = self.extract(root, sheet, "SITE2-0042")
            self.assertIn("SITE2-0042", out)
            self.assertTrue((Path(root) / "Videos" / "SITE2-0042" / "force_data" / "SITE2-0042_points_ABCD_force.xlsx").exists())
            self.assertFalse((Path(root) / "Videos" / "CASE001" / "force_data" / "CASE001_points_ABCD_force.xlsx").exists())
            self.extract(root, sheet)                                                     # no filter: every case
            self.assertTrue((Path(root) / "Videos" / "CASE001" / "force_data" / "CASE001_points_ABCD_force.xlsx").exists())

    def test_missing_columns_are_named(self):
        with tempfile.TemporaryDirectory() as root:
            sheet = self.build(root)
            wb = openpyxl.Workbook()
            wb.active.append(["Study ID", "Point A start", "Point A end"])
            wb.save(Path(root) / "short.xlsx")
            with self.assertRaises(SystemExit) as cm:
                self.extract(root, Path(root) / "short.xlsx")
            self.assertIn("Point B start", str(cm.exception))
            self.assertIn("Point D end", str(cm.exception))

    def test_missing_tracking_sheet(self):
        with tempfile.TemporaryDirectory() as root:
            ex.BASE = Path(root)
            with self.assertRaises(SystemExit):
                self.extract(root, Path(root) / "nope.xlsx")

    def out_path(self, root, case="CASE001"):
        return Path(root) / "Videos" / case / "force_data" / f"{case}_points_ABCD_force.xlsx"

    def test_a_named_case_that_is_not_in_the_sheet_is_an_error(self):
        with tempfile.TemporaryDirectory() as root:
            sheet = self.build(root)
            with self.assertRaises(SystemExit) as cm:
                self.extract(root, sheet, "NOT-A-CASE")
            self.assertEqual(cm.exception.code.split(":")[0], "ERROR")             # a message, so the exit code is 1
            self.assertIn("NOT-A-CASE", cm.exception.code)
            self.assertIn("not in the tracking sheet", cm.exception.code)
            self.assertFalse(self.out_path(root).exists())

    def test_a_named_case_without_a_force_json_is_an_error(self):
        with tempfile.TemporaryDirectory() as root:
            sheet = self.build(root)
            (Path(root) / "Videos" / "CASE001" / "force_data" / "CASE001_force_raw.json").unlink()
            with self.assertRaises(SystemExit) as cm:
                self.extract(root, sheet, "CASE001")
            self.assertIn("force JSON not found", cm.exception.code)
            self.assertIn("parse_force_api", cm.exception.code)
            self.assertFalse(self.out_path(root).exists())

    def test_a_named_case_with_no_points_defined_is_an_error(self):
        with tempfile.TemporaryDirectory() as root:
            self.build(root)
            sheet = openpyxl.Workbook()
            ws = sheet.active
            ws.append(["Study ID"] + [f"Point {p} {side}" for p in "ABCD" for side in ("start", "end")])
            ws.append(["CASE001"] + [None] * 8)
            sheet.save(Path(root) / "empty.xlsx")
            with self.assertRaises(SystemExit) as cm:
                self.extract(root, Path(root) / "empty.xlsx", "CASE001")
            self.assertIn("no points defined", cm.exception.code)

    def test_an_unnamed_run_skips_bad_cases_but_fails_when_nothing_is_written(self):
        with tempfile.TemporaryDirectory() as root:
            sheet = self.build(root, cases=("CASE001", "CASE002"))
            (Path(root) / "Videos" / "CASE002" / "force_data" / "CASE002_force_raw.json").unlink()
            out = self.extract(root, sheet)                                       # one case written: not an error
            self.assertIn("CASE002: force JSON not found, skipping", out)
            self.assertTrue(self.out_path(root).exists())
            (Path(root) / "Videos" / "CASE001" / "force_data" / "CASE001_force_raw.json").unlink()
            with self.assertRaises(SystemExit) as cm:
                self.extract(root, sheet)                                         # now nothing can be written
            self.assertIn("no output was written", cm.exception.code)


def make_case(root, with_role=True, offset=0.0, buffer_s=5.0):
    """Frame table + force workbook (Point A: 5 before, 10 window, 5 after; Point B: 3 window rows)."""
    fdir = Path(root) / "Videos" / CASE / "force_data"
    fdir.mkdir(parents=True)
    idx = synthetic_index()
    video = Path(root) / "video.mp4"
    video.write_bytes(b"0" * 123)
    idx.source_name, idx.source_size = video.name, video.stat().st_size
    idx.save(fl.frame_table_path(fdir, CASE))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    layout = {"A": (8373, 8393, 8378, 8387), "B": (8413, 8416, 8413, 8415)}
    for label, (lo, hi, w0, w1) in layout.items():
        ws = wb.create_sheet(f"Point {label}")
        ws.append(["Timestamp", "Seconds", "USM1", "In Window"] + (["Role"] if with_role else []))
        for x in range(lo, hi):
            role = "buffer-before" if x < w0 else ("window" if x <= w1 else "buffer-after")
            ws.append([fl.fmt_time(x), float(x), 1.0, "YES" if role == "window" else ""] + ([role] if with_role else []))
    if with_role:
        sel = wb.create_sheet("Selection")
        sel.append(["Point", "Start (as entered)", "End (as entered)", "Buffer (s)", "Time offset (s)",
                    "Window rows", "Buffer-before rows", "Buffer-after rows", "Force seconds kept"])
        sel.append(["A", "02:19:38", "02:19:47", buffer_s, offset, 10, 5, 5, "8373 to 8392"])
        sel.append(["B", "02:20:13", "02:20:15", buffer_s, offset, 3, 0, 0, "8413 to 8415"])
    wb.save(fdir / f"{CASE}_points_ABCD_force.xlsx")
    return fdir, video


def run_quiet(func, *args, **kw):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        func(*args, **kw)
    return buf.getvalue()


class TestPlanMeasurements(unittest.TestCase):
    def plan(self, root, *argv, **kw):
        pm.BASE = Path(root)
        out = run_quiet(pm.main, [CASE, *argv], **kw)
        with open(fl.plan_path(Path(root) / "Videos" / CASE / "force_data", CASE)) as f:
            return json.load(f), out

    def test_default_plan_counts_and_parameters(self):
        with tempfile.TemporaryDirectory() as root:
            make_case(root)
            plan, out = self.plan(root, "--yes")
            a, b = plan["points"]["A"], plan["points"]["B"]
            self.assertEqual((a["force_values"], a["window_values"], a["buffer_values"]), (20, 10, 10))
            self.assertEqual(a["max_frames"], 29 * 20)              # 20 contiguous seconds, no frame shared
            self.assertEqual((a["stride"], a["planned_frames"]), (10, 3 * 20))
            self.assertEqual((b["force_values"], b["max_frames"], b["planned_frames"]), (3, 87, 9))
            p = plan["parameters"]
            self.assertEqual((p["half_window_s"], p["half_window_frames"]), (0.467, 14))
            self.assertNotIn("layout", p)
            self.assertEqual((p["time_offset_s"], p["buffer_s"], p["window_center_shift_s"]), (0.0, 5.0, 0.0))
            self.assertEqual(plan["pipeline_version"], fl.PIPELINE_VERSION)
            self.assertIn("Total", out)
            self.assertIn("Max frames", out)
            self.assertIn(str(580 + 87), out)                          # total maximum frames
            self.assertIn("Frames used for two force seconds: 0", out)

    def test_per_point_and_single_stride_flags(self):
        with tempfile.TemporaryDirectory() as root:
            make_case(root)
            plan, _ = self.plan(root, "--stride", "A=5", "B=1")
            self.assertEqual((plan["points"]["A"]["stride"], plan["points"]["A"]["planned_frames"]), (5, 5 * 20))
            self.assertEqual((plan["points"]["B"]["stride"], plan["points"]["B"]["planned_frames"]), (1, 87))
            plan, _ = self.plan(root, "--stride", "7")
            self.assertEqual({p["stride"] for p in plan["points"].values()}, {7})

    def test_interactive_prompt_with_invalid_replies(self):
        with tempfile.TemporaryDirectory() as root:
            make_case(root)
            replies = iter(["x", "0", "", "2"])             # point A: two bad answers then Enter; point B: 2
            plan, out = self.plan(root, interactive=True, input_fn=lambda prompt: next(replies))
            self.assertEqual(plan["points"]["A"]["stride"], 10)
            self.assertEqual(plan["points"]["B"]["stride"], 2)
            self.assertIn("please enter a whole number", out)

    def test_not_a_terminal_uses_default_with_a_note(self):
        with tempfile.TemporaryDirectory() as root:
            make_case(root)
            plan, out = self.plan(root, interactive=False)
            self.assertEqual({p["stride"] for p in plan["points"].values()}, {10})
            self.assertIn("not an interactive terminal", out)

    def test_half_window_in_seconds_follows_the_frame_table(self):
        with tempfile.TemporaryDirectory() as root:
            make_case(root)
            plan, _ = self.plan(root, "--yes", "--half-window-s", "0.4")
            self.assertEqual((plan["parameters"]["half_window_s"], plan["parameters"]["half_window_frames"]), (0.4, 12))
            self.assertEqual(plan["points"]["A"]["max_frames"], 25 * 20)

    def test_half_window_of_half_a_second_or_more_is_refused(self):
        with tempfile.TemporaryDirectory() as root:
            fdir, _ = make_case(root)
            pm.BASE = Path(root)
            for seconds in ("0.5", "1.0"):               # 15 and 30 frames: a frame would serve two force seconds
                with self.assertRaises(SystemExit) as cm:
                    run_quiet(pm.main, [CASE, "--yes", "--half-window-s", seconds])
                self.assertIn("largest allowed is 14 frames", str(cm.exception))
            self.assertFalse(fl.plan_path(fdir, CASE).exists())          # nothing was written

    def test_the_edge_layout_option_is_gone(self):
        with tempfile.TemporaryDirectory() as root:
            make_case(root)
            pm.BASE = Path(root)
            with contextlib.redirect_stderr(io.StringIO()):
                with self.assertRaises(SystemExit) as cm:
                    run_quiet(pm.main, [CASE, "--yes", "--layout", "edge"])
            self.assertEqual(cm.exception.code, 2)                       # argparse: unrecognised argument

    def test_time_offset_and_buffer_come_from_the_selection_sheet(self):
        with tempfile.TemporaryDirectory() as root:
            make_case(root, offset=-1.5, buffer_s=3.0)
            plan, _ = self.plan(root, "--yes")
            self.assertEqual((plan["parameters"]["time_offset_s"], plan["parameters"]["buffer_s"]), (-1.5, 3.0))

    def test_legacy_workbook_counts_window_rows_only_and_says_so(self):
        with tempfile.TemporaryDirectory() as root:
            make_case(root, with_role=False)
            plan, out = self.plan(root, "--yes")
            self.assertEqual(plan["points"]["A"]["force_values"], 10)
            self.assertIn("no Role column", out)

    def test_rerun_replaces_the_plan(self):
        with tempfile.TemporaryDirectory() as root:
            make_case(root)
            self.plan(root, "--yes")
            plan, out = self.plan(root, "--stride", "5")
            self.assertIn("Replacing the plan", out)
            self.assertEqual(plan["points"]["A"]["stride"], 5)

    def test_errors_exit_cleanly(self):
        with tempfile.TemporaryDirectory() as root:
            pm.BASE = Path(root)
            with self.assertRaises(SystemExit):
                pm.main([CASE, "--yes"])                                 # nothing there yet
            fdir, _ = make_case(root)
            for bad in (["--stride", "Z=5"], ["--stride", "A=0"]):
                with self.assertRaises(SystemExit):
                    run_quiet(pm.main, [CASE, *bad])
            fl.frame_table_path(fdir, CASE).unlink()
            with self.assertRaises(SystemExit) as cm:
                run_quiet(pm.main, [CASE, "--yes"])
            self.assertIn("index_video.py", str(cm.exception))

    def test_parse_strides(self):
        self.assertEqual(pm.parse_strides(["10"], ["A", "B"]), {"A": 10, "B": 10})
        self.assertEqual(pm.parse_strides(["A=5", "10"], ["A", "B"]), {"A": 5, "B": 10})
        self.assertEqual(pm.parse_strides(None, ["A"]), {})


class TestMarkerPlanHandling(unittest.TestCase):
    def setup(self, root):
        fdir, video = make_case(root)
        pm.BASE = Path(root)
        run_quiet(pm.main, [CASE, "--yes"])
        ms.BASE = Path(root)
        return fdir, video

    def listing(self, video, *extra):
        return run_quiet(ms.main, [str(video), CASE, "A", "--list", *extra])

    def test_plan_supplies_parameters_and_roles_are_shown(self):
        with tempfile.TemporaryDirectory() as root:
            fdir, video = self.setup(root)
            out = self.listing(video)
            self.assertIn("20 force seconds, 60 distinct frames", out)
            self.assertIn("[buffer-before]", out)
            self.assertIn("[window]", out)
            self.assertIn("[buffer-after]", out)
            self.assertIn("offsets [-10 +0 +10]", out)

    def test_command_line_overrides_the_plan(self):
        with tempfile.TemporaryDirectory() as root:
            fdir, video = self.setup(root)
            out = self.listing(video, "--stride", "5")
            self.assertIn("offsets [-10 -5 +0 +5 +10]", out)
            self.assertIn("100 distinct frames", out)

    def test_plan_wins_over_stored_parameters(self):
        with tempfile.TemporaryDirectory() as root:
            fdir, video = self.setup(root)
            store = fl.MarksStore(fdir / f"{CASE}_frame_marks.json", CASE)
            store.note_params("A", fl.resolve_params(stride=2))
            store.save()
            self.assertIn("offsets [-10 +0 +10]", self.listing(video))

    def test_half_window_over_the_limit_is_refused_with_instructions(self):
        with tempfile.TemporaryDirectory() as root:
            fdir, video = self.setup(root)
            with self.assertRaises(SystemExit) as cm:
                self.listing(video, "--half-window-frames", "15")
            self.assertIn("largest allowed is 14 frames", str(cm.exception))
            self.assertIn("plan_measurements.py", str(cm.exception))

    def test_a_plan_made_by_2_0_0_is_refused_until_it_is_remade(self):
        with tempfile.TemporaryDirectory() as root:
            fdir, video = self.setup(root)
            path = fl.plan_path(fdir, CASE)
            old = json.loads(path.read_text())
            old["pipeline_version"] = "2.0.0"
            old["parameters"].update({"half_window_s": 0.5, "half_window_frames": 15, "layout": "edge"})
            path.write_text(json.dumps(old))
            with self.assertRaises(SystemExit) as cm:
                self.listing(video)
            self.assertIn("largest allowed is 14 frames", str(cm.exception))
            run_quiet(pm.main, [CASE, "--yes"])                         # re-planning fixes it
            self.assertIn("offsets [-10 +0 +10]", self.listing(video))

    def test_missing_plan_or_frame_table_stops_with_instructions(self):
        with tempfile.TemporaryDirectory() as root:
            fdir, video = self.setup(root)
            fl.plan_path(fdir, CASE).unlink()
            with self.assertRaises(SystemExit) as cm:
                self.listing(video)
            self.assertIn("plan_measurements.py", str(cm.exception))
            fl.frame_table_path(fdir, CASE).unlink()
            with self.assertRaises(SystemExit) as cm:
                self.listing(video)
            self.assertIn("index_video.py", str(cm.exception))

    def test_a_different_video_is_refused(self):
        with tempfile.TemporaryDirectory() as root:
            fdir, video = self.setup(root)
            video.write_bytes(b"1" * 999)
            with self.assertRaises(SystemExit):
                self.listing(video)

    def test_session_describe_names_the_role(self):
        idx = synthetic_index()
        store = fl.MarksStore(Path(tempfile.mkdtemp()) / "m.json", CASE)
        plan = fl.plan_point(idx, [8377, 8378], stride=10)
        s = ms.MarkSession(store, "A", plan, 10.0, idx, roles={8377.0: "buffer-before", 8378.0: "window"})
        text = s.describe(plan.frames[0])
        self.assertIn("(buffer-before)", text)


if __name__ == "__main__":
    unittest.main()
