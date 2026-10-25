"""
export_calibration_frame.py (the exact first frame of a point's segment, D2) and the stepping logic of
pick_calibration_frame.py (choosing another frame when the default is unusable, D3), on a tiny synthetic video
made with PyAV (skipped if PyAV is missing). No display is needed.

Run from the repository root:
    python3 -m unittest discover -s tests -v
"""

import contextlib
import io
import json
import os
import sys
import tempfile
import unittest
import warnings
from pathlib import Path

import numpy as np
import openpyxl

os.environ.setdefault("MPLBACKEND", "Agg")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import framelib as fl  # noqa: E402
import index_video as iv  # noqa: E402
from test_index_video import make_video  # noqa: E402

try:
    import av  # noqa: F401
    from PIL import Image
except ImportError:
    av = None
if av:
    import export_calibration_frame as ecf  # noqa: E402
    import pick_calibration_frame as pcf  # noqa: E402

CASE = "case1"


def quiet(func, *args, **kw):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        func(*args, **kw)
    return buf.getvalue()


def build_case(root, start="00:00:01", end="00:00:02", buffer_s=0.5):
    """A 90-frame video (frame k is a flat grey of (3k) % 250), its frame table, and a force workbook whose
    Selection sheet records point A's window (video time 1 s to 2 s) the way extract_point_forces.py does."""
    root = Path(root)
    fdir = root / "Videos" / CASE / "force_data"
    fdir.mkdir(parents=True)
    video = root / "clip.mp4"
    make_video(video)
    iv.BASE = root
    quiet(iv.main, [str(video), CASE])
    wb = openpyxl.Workbook()
    wb.active.title = "Point A"
    wb.active.append(["Timestamp", "Seconds", "USM1", "In Window", "Role"])
    sel = wb.create_sheet("Selection")
    sel.append(["Point", "Start (as entered)", "End (as entered)", "Buffer (s)", "Time offset (s)",
                "Window rows", "Buffer-before rows", "Buffer-after rows", "Force seconds kept"])
    sel.append(["A", start, end, buffer_s, 0.0, 2, 1, 1, "0 to 3"])
    wb.save(fdir / f"{CASE}_points_ABCD_force.xlsx")
    ecf.BASE = pcf.BASE = root
    return video, fdir



def setUpModule():
    # plt.show() under the headless Agg backend warns that it cannot show the window: expected here, not a failure
    warnings.filterwarnings("ignore", message="FigureCanvasAgg is non-interactive")


class TestReferenceCrop(unittest.TestCase):
    """The square of the previous reference frame shown around its two landmark clicks (R1)."""
    SHAPE = (1024, 1280)          # height, width

    def test_square_centred_on_the_midpoint_with_side_three_times_the_landmark_distance(self):
        x0, x1, y0, y1 = fl.reference_crop(self.SHAPE, (600, 500), (700, 500))        # 100 px apart: side 300
        self.assertEqual((x1 - x0, y1 - y0), (300, 300))
        self.assertEqual(((x0 + x1) / 2, (y0 + y1) / 2), (650, 500))

    def test_a_floor_applies_to_close_landmarks(self):
        x0, x1, y0, y1 = fl.reference_crop(self.SHAPE, (600, 500), (630, 500))        # 30 px apart: 3 x = 90 < 300
        self.assertEqual((x1 - x0, y1 - y0), (300, 300))

    def test_the_square_is_shifted_to_stay_inside_the_image(self):
        x0, x1, y0, y1 = fl.reference_crop(self.SHAPE, (20, 30), (80, 90))
        self.assertEqual((x0, y0), (0, 0))
        self.assertEqual((x1 - x0, y1 - y0), (300, 300))
        x0, x1, y0, y1 = fl.reference_crop(self.SHAPE, (1250, 1000), (1270, 1020))
        self.assertEqual((x1, y1), (1280, 1024))
        self.assertEqual((x1 - x0, y1 - y0), (300, 300))

    def test_the_square_never_exceeds_the_image(self):
        x0, x1, y0, y1 = fl.reference_crop((200, 500), (100, 100), (400, 150))        # side would be 912
        self.assertEqual((x0, x1, y0, y1), (150, 350, 0, 200))                        # limited to the short side, kept inside
        self.assertTrue(0 <= x0 < x1 <= 500 and 0 <= y0 < y1 <= 200)

    def test_both_landmarks_are_always_inside_the_crop(self):
        for p1, p2 in (((10, 10), (130, 40)), ((1200, 900), (1100, 980)), ((640, 512), (700, 600))):
            x0, x1, y0, y1 = fl.reference_crop(self.SHAPE, p1, p2)
            for x, y in (p1, p2):
                self.assertTrue(x0 <= x < x1 and y0 <= y < y1, (p1, p2, (x0, x1, y0, y1)))


class TestSegmentHelpers(unittest.TestCase):
    def test_parse_hms(self):
        self.assertEqual(fl.parse_hms("02:19:38"), 8378.0)
        self.assertEqual(fl.parse_hms("00:00:01.250"), 1.25)
        self.assertEqual(fl.parse_hms("12.5"), 12.5)
        with self.assertRaises(ValueError):
            fl.parse_hms("soon")

    def test_position_of_a_time_in_the_segment(self):
        seg = {"start_s": 10.0, "end_s": 20.0, "buffer_s": 5.0}
        self.assertEqual(fl.segment_position(seg, 10.0), "window")
        self.assertEqual(fl.segment_position(seg, 20.0), "window")
        self.assertEqual(fl.segment_position(seg, 9.99), "buffer-before")
        self.assertEqual(fl.segment_position(seg, 5.0), "buffer-before")
        self.assertEqual(fl.segment_position(seg, 20.01), "buffer-after")
        self.assertEqual(fl.segment_position(seg, 4.99), "outside the selected range")
        self.assertEqual(fl.segment_position(seg, 25.01), "outside the selected range")


@unittest.skipUnless(av, "PyAV or Pillow is not installed")
class TestSelectionSheet(unittest.TestCase):
    def test_the_selection_sheet_gives_the_segment_of_each_point(self):
        with tempfile.TemporaryDirectory() as root:
            video, fdir = build_case(root)
            segs = fl.read_selection(fdir / f"{CASE}_points_ABCD_force.xlsx")
            self.assertEqual(set(segs), {"A"})
            self.assertEqual((segs["A"]["start_s"], segs["A"]["end_s"], segs["A"]["buffer_s"]), (1.0, 2.0, 0.5))
            self.assertEqual(segs["A"]["start_text"], "00:00:01")

    def test_a_workbook_without_a_selection_sheet_gives_nothing(self):
        with tempfile.TemporaryDirectory() as root:
            wb = openpyxl.Workbook()
            wb.save(Path(root) / "w.xlsx")
            self.assertEqual(fl.read_selection(Path(root) / "w.xlsx"), {})


@unittest.skipUnless(av, "PyAV or Pillow is not installed")
class TestExportCalibrationFrame(unittest.TestCase):
    def png(self, root, name):
        return Path(root) / "Videos" / CASE / "calibration_frames" / name

    def test_default_is_the_first_frame_of_the_buffer_at_native_resolution(self):
        with tempfile.TemporaryDirectory() as root:
            video, fdir = build_case(root)                                     # window 1 s to 2 s, buffer 0.5 s
            out = quiet(ecf.main, [str(video), CASE, "A"])
            path = self.png(root, f"{CASE}_A_frame15.png")                     # buffer start 0.500 s = frame 15
            self.assertTrue(path.exists(), out)
            img = np.asarray(Image.open(path).convert("RGB"))
            self.assertEqual(img.shape, (48, 64, 3))                           # the video's own size
            idx = fl.load_index(fdir, CASE, video)
            with fl.FrameReader(video, idx) as reader:
                self.assertTrue(np.array_equal(img, reader.get(15)))          # lossless: the decoded frame itself
            info = json.loads(path.with_suffix(".json").read_text())
            self.assertEqual((info["frame"], info["is_default_frame"], info["default_frame"]), (15, True, 15))
            self.assertEqual(info["position"], "buffer-before")
            self.assertEqual(info["crc32"], f"{fl.frame_crc32(img):08x}")
            self.assertEqual((info["case"], info["point"], info["pipeline_version"]),
                             (CASE, "A", fl.PIPELINE_VERSION))
            self.assertIn("first frame of the buffer", out)

    def test_the_default_is_never_before_the_start_of_the_buffer(self):
        # window start 1.010 s, buffer 0.5 s: the buffer starts at 0.510 s, 10 ms after frame 15 (0.500 s) and 23 ms
        # before frame 16 (0.533 s). The nearest frame is 15, which lies outside; the first frame of the buffer is 16.
        with tempfile.TemporaryDirectory() as root:
            video, fdir = build_case(root, start="00:00:01.010")
            quiet(ecf.main, [str(video), CASE, "A"])
            path = self.png(root, f"{CASE}_A_frame16.png")
            self.assertTrue(path.exists())
            info = json.loads(path.with_suffix(".json").read_text())
            self.assertEqual((info["frame"], info["is_default_frame"], info["position"]), (16, True, "buffer-before"))
            self.assertGreaterEqual(info["video_time_s"], 0.510)

    def test_a_buffer_start_exactly_on_a_frame_uses_that_frame(self):
        with tempfile.TemporaryDirectory() as root:
            video, fdir = build_case(root, start="00:00:01")
            quiet(ecf.main, [str(video), CASE, "A"])
            self.assertTrue(self.png(root, f"{CASE}_A_frame15.png").exists())

    def test_without_a_buffer_the_default_is_the_window_start(self):
        with tempfile.TemporaryDirectory() as root:
            video, fdir = build_case(root, buffer_s=0.0)
            quiet(ecf.main, [str(video), CASE, "A"])
            path = self.png(root, f"{CASE}_A_frame30.png")
            info = json.loads(path.with_suffix(".json").read_text())
            self.assertEqual((info["frame"], info["is_default_frame"], info["position"]), (30, True, "window"))

    def test_another_frame_is_recorded_as_a_deviation(self):
        with tempfile.TemporaryDirectory() as root:
            video, fdir = build_case(root)
            out = quiet(ecf.main, [str(video), CASE, "A", "--frame", "20"])
            path = self.png(root, f"{CASE}_A_frame20.png")
            info = json.loads(path.with_suffix(".json").read_text())
            self.assertEqual((info["frame"], info["is_default_frame"], info["default_frame"]), (20, False, 15))
            self.assertEqual(info["position"], "buffer-before")
            self.assertAlmostEqual(info["seconds_from_window_start"], -0.333, places=2)
            self.assertIn("not the default", out)
            self.assertFalse(self.png(root, f"{CASE}_A_frame30.png").exists())       # only what was asked for

    def test_the_sidecar_sits_next_to_the_png_under_the_name_calibrate_reference_reads(self):
        # calibrate_reference.read_frame_info(png) opens png.with_suffix(".json") (see test_calibration.py);
        # it is not imported here because OpenCV and PyAV do not share a process well on macOS.
        with tempfile.TemporaryDirectory() as root:
            video, fdir = build_case(root)
            quiet(ecf.main, [str(video), CASE, "A", "--frame", "20"])
            side = self.png(root, f"{CASE}_A_frame20.png").with_suffix(".json")
            info = json.loads(side.read_text())
            self.assertEqual((info["frame"], info["is_default_frame"]), (20, False))

    def test_explicit_output_path(self):
        with tempfile.TemporaryDirectory() as root:
            video, fdir = build_case(root)
            target = Path(root) / "mine" / "still.png"
            quiet(ecf.main, [str(video), CASE, "A", "--out", str(target)])
            self.assertTrue(target.exists())
            self.assertTrue(target.with_suffix(".json").exists())

    def test_the_ruler_frame_is_exported_by_frame_number_or_time(self):
        with tempfile.TemporaryDirectory() as root:
            video, fdir = build_case(root)
            quiet(ecf.main, [str(video), CASE, "ruler", "--frame", "40"])
            path = self.png(root, f"{CASE}_ruler_frame40.png")
            info = json.loads(path.with_suffix(".json").read_text())
            self.assertEqual((info["point"], info["frame"], info["is_default_frame"]), ("ruler", 40, None))
            idx = fl.load_index(fdir, CASE, video)
            with fl.FrameReader(video, idx) as reader:
                self.assertTrue(np.array_equal(np.asarray(Image.open(path).convert("RGB")), reader.get(40)))
            out = quiet(ecf.main, [str(video), CASE, "RULER", "--time", "00:00:01.5"])      # first frame at or after 1.5 s
            self.assertTrue(self.png(root, f"{CASE}_ruler_frame45.png").exists(), out)

    def test_the_ruler_frame_needs_a_frame_or_a_time_and_a_point_needs_neither_a_time(self):
        with tempfile.TemporaryDirectory() as root:
            video, fdir = build_case(root)
            with self.assertRaises(SystemExit) as cm:
                quiet(ecf.main, [str(video), CASE, "ruler"])
            self.assertIn("--frame", str(cm.exception.code))
            with self.assertRaises(SystemExit) as cm:
                quiet(ecf.main, [str(video), CASE, "A", "--time", "00:00:01"])
            self.assertIn("ruler", str(cm.exception.code))

    def test_clear_errors(self):
        with tempfile.TemporaryDirectory() as root:
            video, fdir = build_case(root)
            for argv, expect in (([str(video), CASE, "A", "--frame", "500"], "outside the video"),
                                 ([str(video), CASE, "A", "--frame", "-1"], "outside the video"),
                                 ([str(video), CASE, "C"], "extract_point_forces.py")):
                with self.subTest(argv=argv):
                    with self.assertRaises(SystemExit) as cm:
                        quiet(ecf.main, argv)
                    self.assertIn(expect, str(cm.exception.code))
            (fdir / f"{CASE}_points_ABCD_force.xlsx").unlink()
            with self.assertRaises(SystemExit) as cm:
                quiet(ecf.main, [str(video), CASE, "A"])
            self.assertIn("extract_point_forces.py", str(cm.exception.code))
            fl.frame_table_path(fdir, CASE).unlink()
            with self.assertRaises(SystemExit) as cm:
                quiet(ecf.main, [str(video), CASE, "A"])
            self.assertIn("index_video.py", str(cm.exception.code))


@unittest.skipUnless(av, "PyAV or Pillow is not installed")
class TestFramePicker(unittest.TestCase):
    """The stepping logic behind the scrubbing window."""

    def make(self, root, **kw):
        video, fdir = build_case(root)
        idx = fl.load_index(fdir, CASE, video)
        seg = fl.read_selection(fdir / f"{CASE}_points_ABCD_force.xlsx")["A"]
        return pcf.FramePicker(idx, seg, **kw), idx

    def test_it_opens_on_the_first_frame_of_the_buffer(self):
        with tempfile.TemporaryDirectory() as root:
            p, idx = self.make(root)
            self.assertEqual((p.frame, p.default_frame), (15, 15))
            self.assertTrue(p.on_default())
            self.assertEqual((p.lo, p.hi), (15, 75))                             # 0.5 s buffer before 1 s and after 2 s

    def test_the_range_ends_are_inside_the_buffered_window(self):
        # buffer start 0.510 s (between frames 15 at 0.500 and 16 at 0.533) and buffer end 2.525 s (between frames 75 at
        # 2.500 and 76 at 2.533): the nearest frames 15 and 76 are outside, the range is 16 to 75 (issue 17).
        with tempfile.TemporaryDirectory() as root:
            video, fdir = build_case(root, start="00:00:01.010", end="00:00:02.025")
            idx = fl.load_index(fdir, CASE, video)
            seg = fl.read_selection(fdir / f"{CASE}_points_ABCD_force.xlsx")["A"]
            p = pcf.FramePicker(idx, seg)
            self.assertEqual((p.lo, p.hi), (16, 75))
            for k in (p.lo, p.hi):
                self.assertNotEqual(fl.segment_position(seg, float(idx.pts_s[k])), "outside the selected range")
            self.assertEqual(p.default_frame, 16)

    def test_steps_clamp_to_the_selected_range(self):
        with tempfile.TemporaryDirectory() as root:
            p, _ = self.make(root)
            p.step(+1)
            self.assertEqual(p.frame, 16)
            p.step(+30)
            self.assertEqual(p.frame, 46)
            p.step(-100)
            self.assertEqual(p.frame, 15)                                        # not before the buffer
            p.step(+1000)
            self.assertEqual(p.frame, 75)

    def test_seconds_use_the_stored_frame_times(self):
        with tempfile.TemporaryDirectory() as root:
            p, idx = self.make(root)
            p.step_seconds(+1)
            self.assertEqual(p.frame, 45)                                        # 1.500 s
            p.step_seconds(-1)
            self.assertEqual(p.frame, 15)

    def test_default_key_returns_to_the_start_of_the_buffer(self):
        with tempfile.TemporaryDirectory() as root:
            p, _ = self.make(root)
            p.step(+25)
            self.assertFalse(p.on_default())
            p.to_default()
            self.assertEqual(p.frame, 15)

    def test_label_names_time_frame_and_position(self):
        with tempfile.TemporaryDirectory() as root:
            p, idx = self.make(root)
            self.assertIn("frame 15", p.label())
            self.assertIn("00:00:00.500", p.label())
            self.assertIn("buffer-before", p.label())
            self.assertIn("start of the buffer", p.label())
            p.step(+15)
            self.assertIn("window", p.label())
            self.assertIn("at the window start", p.label())
            p.step(+10)
            self.assertIn("0.33 s after the window start", p.label())
            self.assertNotIn("start of the buffer", p.label())

    def test_keys_move_the_frame_and_enter_exports_it(self):
        with tempfile.TemporaryDirectory() as root:
            video, fdir = build_case(root)
            idx = fl.load_index(fdir, CASE, video)
            seg = fl.read_selection(fdir / f"{CASE}_points_ABCD_force.xlsx")["A"]
            picker = pcf.FramePicker(idx, seg)
            exported = []
            with fl.FrameReader(video, idx) as reader:
                fig, on_key, state = pcf.make_window(picker, reader, lambda k: exported.append(k) or f"frame{k}.png")
                for key, frame in (("right", 16), ("shift+right", 26), ("up", 56), ("down", 26), ("d", 15),
                                   ("right", 16), ("x", 16)):                       # an unknown key changes nothing
                    on_key(type("E", (), {"key": key})())
                    self.assertEqual(picker.frame, frame, key)
                on_key(type("E", (), {"key": "enter"})())
            self.assertEqual((exported, state["exported"], state["quit"]), ([16], "frame16.png", False))

    def test_the_previous_points_calibration_frame_is_shown_zoomed_beside_the_current_frame(self):
        with tempfile.TemporaryDirectory() as root:
            video, fdir = build_case(root)
            idx = fl.load_index(fdir, CASE, video)
            seg = fl.read_selection(fdir / f"{CASE}_points_ABCD_force.xlsx")["A"]
            prev_png = Path(root) / "prev.png"
            Image.fromarray(np.full((480, 640, 3), 90, np.uint8)).save(prev_png)
            ref = {"name": "the ruler frame", "image": str(prev_png), "p1": (300, 200), "p2": (380, 260), "missing": None}
            with fl.FrameReader(video, idx) as reader:
                fig, on_key, state = pcf.make_window(pcf.FramePicker(idx, seg), reader, lambda k: "x", previous=ref)
                titles = [a.get_title() for a in fig.axes]
                fig2, _, _ = pcf.make_window(pcf.FramePicker(idx, seg), reader, lambda k: "x")        # no previous frame
            self.assertTrue(any("the ruler frame" in s for s in titles), titles)
            self.assertEqual(len(fig.axes), len(fig2.axes) + 2)             # the zoomed panel and its whole-frame thumbnail
            self.assertTrue(any("Default frame" in s for s in titles), titles)

    def test_q_quits_without_exporting(self):
        with tempfile.TemporaryDirectory() as root:
            video, fdir = build_case(root)
            idx = fl.load_index(fdir, CASE, video)
            seg = fl.read_selection(fdir / f"{CASE}_points_ABCD_force.xlsx")["A"]
            with fl.FrameReader(video, idx) as reader:
                fig, on_key, state = pcf.make_window(pcf.FramePicker(idx, seg), reader, lambda k: "never")
                on_key(type("E", (), {"key": "q"})())
            self.assertEqual((state["exported"], state["quit"]), (None, True))

    def test_closing_the_window_without_exporting_is_an_error(self):
        # no display here, so plt.show() returns at once: the same situation as closing the window early
        with tempfile.TemporaryDirectory() as root:
            video, fdir = build_case(root)
            with self.assertRaises(SystemExit) as cm:
                with contextlib.redirect_stdout(io.StringIO()):
                    pcf.main([str(video), CASE, "A"])
            self.assertIn("No frame exported", str(cm.exception.code))
            self.assertFalse((Path(root) / "Videos" / CASE / "calibration_frames").exists())

    def test_a_start_frame_outside_the_range_is_pulled_in(self):
        with tempfile.TemporaryDirectory() as root:
            p, _ = self.make(root, start_frame=2)
            self.assertEqual(p.frame, 15)
            self.assertEqual(pcf.FramePicker(_, p.seg, 500).frame, 75)


if __name__ == "__main__":
    unittest.main()
