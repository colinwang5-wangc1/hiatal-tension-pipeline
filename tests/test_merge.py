"""
Tests for merge_distance_force.py on a synthetic case with a known answer: the gap
closes at exactly 5 mm/s (d = 30 - 5 * (t - 8378) mm), so the straight-line fit at the
sample time, the slope and the SD can be checked against exact values.

Point A has 5 buffer-before rows, 10 window rows and 5 buffer-after rows.

Run from the repository root:
    python3 -m unittest discover -s tests -v
"""

import csv
import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

import numpy as np
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import framelib as fl  # noqa: E402
import merge_distance_force as mg  # noqa: E402
from test_framelib import synthetic_index  # noqa: E402

CASE = "case1"
SLOPE = -5.0          # mm per second
PXMM = 10.0
SECONDS = list(range(8373, 8393))          # 5 buffer-before, 10 window (8378-8387), 5 buffer-after


def true_d(t):
    return 30.0 + SLOPE * (t - 8378.0)


def role_of(x):
    return "buffer-before" if x < 8378 else ("window" if x <= 8387 else "buffer-after")


SUMMARY = "Summary A (from Frames)"


def build_case(root, idx, params, with_role=True, measured=SECONDS, calibration_frame=None):
    """Force workbook, frame table and marks. `measured`: force seconds whose frames are marked."""
    fdir = Path(root) / "Videos" / CASE / "force_data"
    fdir.mkdir(parents=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Point A"
    ws.append(["Timestamp", "Seconds", "USM1", "USM3", "In Window"] + (["Role"] if with_role else []))
    for x in SECONDS:
        ws.append([fl.fmt_time(x), float(x), 1.0 + (x - 8373) * 0.1, None if x % 2 else 0.5,
                   "YES" if role_of(x) == "window" else ""] + ([role_of(x)] if with_role else []))
    wb.save(fdir / f"{CASE}_points_ABCD_force.xlsx")
    idx.save(fl.frame_table_path(fdir, CASE))
    store = fl.MarksStore(fdir / f"{CASE}_frame_marks.json", CASE)
    store.set_video(idx)
    store.note_params("A", params)
    plan = fl.plan_point(idx, measured, **params)
    for k in plan.frames:
        d = true_d(idx.pts_s[k])
        store.put("A", k, {"pts_s": float(idx.pts_s[k]), "left_crus_px": [0, 0], "right_crus_px": [d * PXMM, 0],
                           "distance_px": d * PXMM, "distance_mm": d, "px_per_mm": PXMM, "skipped": False,
                           "source": "manual", "calibration_frame": calibration_frame})
    store.save()
    return fdir, store, plan


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_csv(fdir):
    with open(fdir / f"{CASE}_frame_distances.csv", newline="") as f:
        return list(csv.DictReader(f))


def read_merged(fdir):
    ws = openpyxl.load_workbook(fdir / f"{CASE}_points_ABCD_force_merged.xlsx")[SUMMARY]
    head = [c.value for c in ws[1]]
    return [dict(zip(head, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]


FRAME_COLUMNS = ["case_id", "point", "role", "force_second", "force_time_hms", "frame", "pts", "time_s", "time_hms",
                 "offset_frames", "offset_ms", "status", "skip_reason", "distance_mm", "distance_px",
                 "left_x", "left_y", "right_x", "right_y", "px_per_mm", "calibration_frame", "scale_vs_calibration",
                 "USM1", "USM3", "Total Force (N)"]


def read_frames_sheet(fdir):
    ws = openpyxl.load_workbook(fdir / f"{CASE}_points_ABCD_force_merged.xlsx")["Frames"]
    head = [c.value for c in ws[1]]
    return head, [dict(zip(head, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]


class TestFramesTable(unittest.TestCase):
    """The per-frame table is the product: one row per planned frame, with its distance and the force paired with it."""

    @classmethod
    def setUpClass(cls):
        cls.idx = synthetic_index()

    def run_merge(self, root, *extra):
        mg.BASE = Path(root)
        mg.main([CASE, *extra])

    def test_columns_and_row_count(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir, store, plan = build_case(root, self.idx, params)
            self.run_merge(root)
            head, rows = read_frames_sheet(fdir)
            self.assertEqual(head, FRAME_COLUMNS)
            self.assertEqual(len(rows), 20 * 3)
            csv_rows = read_csv(fdir)
            self.assertEqual(list(csv_rows[0].keys()), FRAME_COLUMNS)

    def test_each_row_carries_its_own_time_and_force(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir, store, plan = build_case(root, self.idx, params, calibration_frame=251300)
            self.run_merge(root)
            _, rows = read_frames_sheet(fdir)
            for r in rows:
                k = r["frame"]
                self.assertEqual(r["case_id"], CASE)
                self.assertEqual(r["pts"], int(self.idx.pts[k]))                          # the stored timestamp, unchanged
                self.assertAlmostEqual(r["time_s"], (int(self.idx.pts[k]) - int(self.idx.pts[0])) / 153600, places=6)
                self.assertEqual(r["time_hms"], fl.fmt_time(self.idx.pts_s[k]))
                x = r["force_second"]
                self.assertEqual(r["force_time_hms"], fl.fmt_time(x)[:8])                  # HH:MM:SS of the force second
                usm1 = 1.0 + (x - 8373) * 0.1
                usm3 = None if int(x) % 2 else 0.5
                self.assertAlmostEqual(r["USM1"], usm1, places=6)
                self.assertEqual(r["USM3"], usm3)                                         # a blank stays blank
                self.assertAlmostEqual(r["Total Force (N)"], usm1 + (usm3 or 0.0), places=6)   # a blank counts as 0 N
                self.assertEqual((r["status"], r["skip_reason"]), ("marked", None))
                self.assertEqual((r["px_per_mm"], r["calibration_frame"]), (PXMM, 251300))

    def test_time_recomputes_from_pts_and_the_info_sheet(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir, store, plan = build_case(root, self.idx, params)
            self.run_merge(root)
            wb = openpyxl.load_workbook(fdir / f"{CASE}_points_ABCD_force_merged.xlsx")
            text = " ".join(str(c.value) for row in wb["Merge info"].iter_rows() for c in row if c.value is not None)
            self.assertIn("Time base 1/153600", text)
            self.assertIn(f"first frame pts {int(self.idx.pts[0])}", text)
            _, rows = read_frames_sheet(fdir)
            self.assertTrue(all(abs(r["time_s"] - (r["pts"] - int(self.idx.pts[0])) * (1 / 153600)) < 1e-6 for r in rows))

    def test_the_sheet_and_the_csv_hold_the_same_table(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir, store, plan = build_case(root, self.idx, params)
            self.run_merge(root)
            _, sheet = read_frames_sheet(fdir)
            csv_rows = read_csv(fdir)
            self.assertEqual(len(sheet), len(csv_rows))
            for s, c in zip(sheet, csv_rows):
                for key, value in s.items():
                    if value is None:
                        self.assertEqual(c[key], "", key)
                    elif isinstance(value, (int, float)):
                        self.assertAlmostEqual(float(c[key]), float(value), places=6, msg=key)
                    else:
                        self.assertEqual(c[key], value, key)

    def test_status_skip_reason_and_scale(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir, store, plan = build_case(root, self.idx, params)
            _, f8380 = fl.sample_frames(self.idx, 8380, **params)
            store.put("A", f8380[0].frame, {"skipped": True, "pts_s": 0.0, "skip_reason": "camera: scale 0.76 vs calibration"})
            store.put("A", f8380[1].frame, {"skipped": True, "pts_s": 0.0})                 # no reason given
            store.delete("A", f8380[2].frame)                                               # left unmarked
            store.set_scale("A", f8380[0].frame, 0.76, 33)
            store.set_scale("A", f8380[2].frame, None, 4)                                   # estimate failed
            store.save()
            self.run_merge(root)
            _, rows = read_frames_sheet(fdir)
            by_frame = {r["frame"]: r for r in rows}
            a, b, c = (by_frame[sf.frame] for sf in f8380)
            self.assertEqual((a["status"], a["skip_reason"], a["scale_vs_calibration"]),
                             ("skipped", "camera: scale 0.76 vs calibration", 0.76))
            self.assertEqual((b["status"], b["skip_reason"], b["distance_mm"], b["px_per_mm"]), ("skipped", None, None, None))
            self.assertEqual((c["status"], c["distance_mm"], c["scale_vs_calibration"]), ("not marked", None, None))
            self.assertEqual({r["status"] for r in rows}, {"marked", "skipped", "not marked"})

    def test_arms_follow_the_force_workbook(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir, store, plan = build_case(root, self.idx, params)
            path = fdir / f"{CASE}_points_ABCD_force.xlsx"
            wb = openpyxl.load_workbook(path)
            ws = wb["Point A"]
            ws.cell(row=1, column=3, value="USM2")                                          # arms USM2 and USM3 instead
            wb.save(path)
            self.run_merge(root)
            head, rows = read_frames_sheet(fdir)
            self.assertEqual(head[-3:], ["USM2", "USM3", "Total Force (N)"])
            self.assertNotIn("USM1", head)

    def test_the_summary_is_labelled_as_derived_and_follows_the_frames_sheet(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir, store, plan = build_case(root, self.idx, params)
            self.run_merge(root)
            wb = openpyxl.load_workbook(fdir / f"{CASE}_points_ABCD_force_merged.xlsx")
            self.assertEqual(wb.sheetnames[0], "Frames")
            self.assertIn(SUMMARY, wb.sheetnames)
            self.assertNotIn("Point A", wb.sheetnames)
            info = " ".join(str(c.value) for row in wb["Merge info"].iter_rows() for c in row if c.value is not None)
            self.assertIn("derived from the Frames sheet", info)
            _, frames = read_frames_sheet(fdir)
            summary = {r["Seconds"]: r for r in read_merged(fdir)}
            for x in (8380.0, 8385.0):
                d = [r["distance_mm"] for r in frames if r["force_second"] == x]
                self.assertAlmostEqual(summary[x]["Distance (mm)"], sum(d) / len(d), places=3)


class TestMerge(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.idx = synthetic_index()

    def run_merge(self, root, *extra):
        mg.BASE = Path(root)
        mg.main([CASE, *extra])

    def test_statistics_match_the_known_linear_gap(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir, store, plan = build_case(root, self.idx, params)
            before = sha(fdir / f"{CASE}_points_ABCD_force.xlsx")
            self.run_merge(root)
            self.assertEqual(sha(fdir / f"{CASE}_points_ABCD_force.xlsx"), before)     # input untouched
            rows = {r["Seconds"]: r for r in read_merged(fdir)}
            r = rows[8380.0]
            _, frames = fl.sample_frames(self.idx, 8380, **params)
            t_rel = np.array([f.offset_ms / 1000 for f in frames])
            self.assertEqual((r["N frames"], r["Frames planned"]), (3, 3))
            self.assertAlmostEqual(r["Fit at x (mm)"], true_d(8380.0), places=3)
            self.assertAlmostEqual(r["Slope (mm/s)"], SLOPE, places=3)
            self.assertLess(r["Fit resid SD (mm)"], 1e-3)
            self.assertAlmostEqual(r["Distance SD (mm)"], abs(SLOPE) * t_rel.std(ddof=1), places=3)
            self.assertAlmostEqual(r["Distance (mm)"], true_d(8380.0) + SLOPE * t_rel.mean(), places=3)
            self.assertGreater(r["Distance max (mm)"], r["Distance (mm)"])
            self.assertLess(r["Distance min (mm)"], r["Distance (mm)"])
            self.assertEqual(r["Total Force (N)"], 1.0 + 7 * 0.1 + (0.5 if 8380 % 2 == 0 else 0))
            self.assertEqual(r["Role"], "window")

    def test_buffer_rows_are_measured_and_labelled(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir, store, plan = build_case(root, self.idx, params)
            self.run_merge(root)
            rows = {r["Seconds"]: r for r in read_merged(fdir)}
            self.assertEqual(rows[8374.0]["Role"], "buffer-before")
            self.assertEqual(rows[8390.0]["Role"], "buffer-after")
            self.assertAlmostEqual(rows[8374.0]["Fit at x (mm)"], true_d(8374.0), places=3)
            self.assertAlmostEqual(rows[8390.0]["Fit at x (mm)"], true_d(8390.0), places=3)
            self.assertEqual(rows[8374.0]["N frames"], 3)
            roles = {r["role"] for r in read_csv(fdir)}
            self.assertEqual(roles, {"buffer-before", "window", "buffer-after"})

    def test_csv_has_every_planned_frame_once_and_none_serves_two_force_seconds(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir, store, plan = build_case(root, self.idx, params)
            self.run_merge(root)
            rows = read_csv(fdir)
            self.assertEqual(len(rows), 20 * 3)                    # 20 seconds x 3 frames
            self.assertEqual({r["status"] for r in rows}, {"marked"})
            frames = [r["frame"] for r in rows]
            self.assertEqual(len(set(frames)), len(frames))        # no frame is listed for two force seconds
            self.assertEqual(rows[0]["role"], "buffer-before")
            self.assertIn("time_s", rows[0])

    def test_merge_info_sheet_has_version_and_counts(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir, store, plan = build_case(root, self.idx, params)
            self.run_merge(root)
            info = openpyxl.load_workbook(fdir / f"{CASE}_points_ABCD_force_merged.xlsx")["Merge info"]
            text = " ".join(str(c.value) for row in info.iter_rows() for c in row if c.value is not None)
            self.assertIn(f"Pipeline version {fl.PIPELINE_VERSION}", text)
            head = [c.value for c in info[2]]
            line = dict(zip(head, [c.value for c in info[3]]))
            self.assertEqual((line["Window seconds with data"], line["Buffer seconds with data"]), (10, 10))

    def test_skipped_and_partial_samples(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir, store, plan = build_case(root, self.idx, params)
            _, f8381 = fl.sample_frames(self.idx, 8381, **params)
            _, f8383 = fl.sample_frames(self.idx, 8383, **params)
            _, f8385 = fl.sample_frames(self.idx, 8385, **params)
            for sf in f8381:                                     # every frame occluded
                store.put("A", sf.frame, {"skipped": True, "pts_s": 0.0})
            for sf in f8383[1:]:                                 # only the first frame kept
                store.put("A", sf.frame, {"skipped": True, "pts_s": 0.0})
            store.put("A", f8385[0].frame, {"skipped": True, "pts_s": 0.0})
            store.save()
            self.run_merge(root)
            rows = {r["Seconds"]: r for r in read_merged(fdir)}
            self.assertIsNone(rows[8381.0]["Distance (mm)"])
            self.assertIsNone(rows[8381.0]["N frames"])
            self.assertEqual(rows[8383.0]["N frames"], 1)
            self.assertIsNone(rows[8383.0]["Distance SD (mm)"])
            self.assertIsNone(rows[8383.0]["Fit at x (mm)"])
            self.assertEqual(rows[8385.0]["N frames"], 2)
            self.assertIsNotNone(rows[8385.0]["Distance SD (mm)"])
            self.assertIsNone(rows[8385.0]["Fit resid SD (mm)"])        # the fit needs 3 frames
            self.assertIn("skipped", [r["status"] for r in read_csv(fdir)])

    def test_resummarise_under_a_denser_stride_counts_unmarked(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir, store, plan = build_case(root, self.idx, params)
            self.run_merge(root, "--stride", "5")
            rows = {r["Seconds"]: r for r in read_merged(fdir)}
            self.assertEqual((rows[8380.0]["N frames"], rows[8380.0]["Frames planned"]), (3, 5))
            self.assertIn("not marked", [r["status"] for r in read_csv(fdir)])

    def test_a_half_window_over_the_limit_is_refused(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            build_case(root, self.idx, params)
            with self.assertRaises(SystemExit) as cm:
                self.run_merge(root, "--half-window-frames", "15")
            self.assertIn("largest allowed is 14 frames", str(cm.exception))

    def test_marks_stamped_by_another_major_version_are_refused_until_the_stamp_is_changed_by_hand(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir, store, plan = build_case(root, self.idx, params)
            path = fdir / f"{CASE}_frame_marks.json"
            data = json.loads(path.read_text())
            data["pipeline_version"] = "2.1.0"
            path.write_text(json.dumps(data))
            with self.assertRaises(SystemExit) as cm:
                self.run_merge(root)
            self.assertIn("2.1.0", str(cm.exception.code))
            self.assertIn("major version", str(cm.exception.code))
            # the per-frame records are unchanged between 2.x and 3.x: the stamp can be edited by hand to reuse them
            data["pipeline_version"] = fl.PIPELINE_VERSION
            path.write_text(json.dumps(data))
            self.run_merge(root)
            _, rows = read_frames_sheet(fdir)
            self.assertEqual({r["status"] for r in rows}, {"marked"})
            self.assertTrue(all(r["scale_vs_calibration"] is None and r["calibration_frame"] is None for r in rows))

    def test_a_stale_2_0_0_half_window_in_the_marks_is_still_refused(self):
        params = fl.resolve_params(stride=10)
        with tempfile.TemporaryDirectory() as root:
            fdir, store, plan = build_case(root, self.idx, params)
            path = fdir / f"{CASE}_frame_marks.json"
            data = json.loads(path.read_text())
            data["points"]["A"]["params_history"][0].update({"half_window_frames": 15})
            path.write_text(json.dumps(data))
            with self.assertRaises(SystemExit) as cm:
                self.run_merge(root)
            self.assertIn("point A", str(cm.exception.code))
            self.assertIn("plan_measurements.py", str(cm.exception.code))
            self.run_merge(root, "--half-window-frames", "14")          # an explicit override still works

    def test_legacy_workbook_without_role_measures_window_rows_only(self):
        params = fl.resolve_params(stride=10)
        window = [x for x in SECONDS if role_of(x) == "window"]
        with tempfile.TemporaryDirectory() as root:
            fdir, store, plan = build_case(root, self.idx, params, with_role=False, measured=window)
            self.run_merge(root)
            rows = {r["Seconds"]: r for r in read_merged(fdir)}
            self.assertIsNotNone(rows[8380.0]["Distance (mm)"])
            self.assertIsNone(rows[8374.0]["Distance (mm)"])
            self.assertEqual({r["role"] for r in read_csv(fdir)}, {"window"})

    def test_missing_inputs_exit_cleanly(self):
        with tempfile.TemporaryDirectory() as root:
            mg.BASE = Path(root)
            with self.assertRaises(SystemExit):
                mg.main([CASE])


if __name__ == "__main__":
    unittest.main()
