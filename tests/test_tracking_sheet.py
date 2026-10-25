"""
make_tracking_sheet.py: creates the tracking spreadsheet with the exact column headers the pipeline needs, and
never overwrites an existing one.

Run from the repository root:
    python3 -m unittest discover -s tests -v
"""

import contextlib
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import extract_point_forces as ex  # noqa: E402
import make_tracking_sheet as mts  # noqa: E402


def quiet(func, *args, **kw):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        func(*args, **kw)
    return buf.getvalue()


class TestTrackingSheet(unittest.TestCase):
    def test_headers_and_one_row_per_case(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "sheet.xlsx"
            quiet(mts.main, ["CASE001", "CASE002", "--out", str(out)])
            ws = openpyxl.load_workbook(out).active
            head = [c.value for c in ws[1]]
            self.assertEqual(head, ["Study ID"] + [f"Point {p} {s}" for p in "ABCD" for s in ("start", "end")])
            self.assertEqual([r[0].value for r in ws.iter_rows(min_row=2)], ["CASE001", "CASE002"])
            self.assertTrue(all(c.value is None for r in ws.iter_rows(min_row=2) for c in r[1:]))

    def test_the_extraction_accepts_the_sheet_once_times_are_filled_in(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "sheet.xlsx"
            quiet(mts.main, ["CASE001", "--out", str(out)])
            wb = openpyxl.load_workbook(out)
            wb.active["B2"], wb.active["C2"] = "00:01:40", "00:01:45"
            wb.save(out)
            fd = Path(d) / "Videos" / "CASE001" / "force_data"
            fd.mkdir(parents=True)
            usm1 = [{"x": x, "y": 1.0, "z": None} for x in range(90, 120)]
            (fd / "CASE001_force_raw.json").write_text(json.dumps({"USM1": usm1, "USM2": [], "USM3": [], "USM4": []}))
            ex.BASE = Path(d)
            quiet(ex.main, ["CASE001", "--tracking-sheet", str(out)])
            self.assertTrue((fd / "CASE001_points_ABCD_force.xlsx").exists())

    def test_an_existing_file_is_never_overwritten(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "sheet.xlsx"
            out.write_text("patient data")
            with self.assertRaises(SystemExit) as cm:
                quiet(mts.main, ["CASE001", "--out", str(out)])
            self.assertIn("already exists", str(cm.exception.code))
            self.assertEqual(out.read_text(), "patient data")

    def test_default_location_is_the_repository_folder(self):
        with tempfile.TemporaryDirectory() as d:
            mts.BASE = Path(d)
            quiet(mts.main, ["CASE001"])
            self.assertTrue((Path(d) / "Hiatal tension patients.xlsx").exists())


if __name__ == "__main__":
    unittest.main()
