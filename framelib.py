"""
framelib.py — Frame identity, sampling windows, exact decoding and mark storage
for the per-frame crural-gap workflow (used by mark_suture_entry.py and
merge_distance_force.py).

Why this exists
---------------
A force sample stamped x is not tied to one video frame. Real frame timestamps
(PTS) are not on an exact nominal-fps grid, and the instant within the second
that a sample represents is not known. Instead of one frame per force second,
the workflow measures several frames around x and reports the spread.

Conventions
-----------
* Frame identity is the raw video's presentation order: frame index k is the
  k-th frame by PTS. Everything is keyed by k and its PTS, never by index / fps.
* Video time zero is the presentation time of the first frame. Any other
  alignment between the force log and the video goes in time_offset_s
  (video time minus force time).
* The window around a force second is built in frames, not milliseconds: the
  centre frame is the frame whose PTS is nearest to
  x + time_offset_s + window_center_shift_s, and the sampled frames are centre
  frame + offset for each offset in the layout. The true millisecond offset of
  every frame is recorded. (A fixed +/-500 ms test on PTS would drop or keep a
  boundary frame depending on where the centre frame falls inside its 33 ms slot.)

Offsets (half_window = H frames, stride = s)
    0, +/-s, +/-2s, ... (<= H): centred on the centre frame.
With H = 14 and s = 10 the offsets are -10, 0, +10; s = 5 gives -10..+10 by 5; s = 1 gives
-14..+14. A denser stride only adds frames.

No frame may serve two force seconds. Force samples are one second apart, which is about 30
frames at 30 fps, and a frame exactly half that far from the centre frame is halfway to the
neighbouring sample. Every offset therefore stays strictly less than half the spacing of the
force samples (max_half_window_frames: 14 at 30 and 29.97 fps, 29 at 60 fps), so one distance
measurement feeds exactly one point of the force-strain curve. sample_frames refuses a larger
half-window, and the default half-window is the largest allowed. (Versions 2.0.0 and earlier
used H = 15 with an "edge" layout whose offsets -15 and +15 were shared by neighbouring seconds.)

Decoding needs PyAV (pip install av). Index building and selection do not.
"""

import datetime
import json
import os
import zlib
from collections import OrderedDict
from dataclasses import dataclass
from pathlib import Path

import numpy as np

PARAM_DEFAULTS = {
    "half_window_frames": None,        # None = the largest allowed for the video (max_half_window_frames)
    "stride": 10,
    "time_offset_s": 0.0,
    "window_center_shift_s": 0.0,
}
SCHEMA_VERSION = 1
PLAN_SCHEMA_VERSION = 1
ROLES = ("window", "buffer-before", "buffer-after")


def _read_version():
    path = Path(__file__).resolve().parent / "VERSION"
    try:
        return path.read_text().strip()
    except OSError:
        raise SystemExit(f"ERROR: the VERSION file is missing next to framelib.py ({path}). Copy the whole folder, not just "
                         "the .py files: every output records the version of the code that made it, and a copy without "
                         "VERSION would silently record 'unknown'.") from None


PIPELINE_VERSION = _read_version()


def check_major_version(stored, what):
    """Refuse files written by a different major version of the pipeline."""
    if stored is None:
        return
    if str(stored).split(".")[0] != PIPELINE_VERSION.split(".")[0]:
        raise SystemExit(f"{what} was written by pipeline version {stored}; this is version {PIPELINE_VERSION}. "
                         "Files from a different major version are not read (see CHANGELOG.md).")


# ---------------------------------------------------------------- time ---------

def fmt_time(t):
    """HH:MM:SS.mmm from seconds, using integer milliseconds so .9995 never prints as 60.000."""
    total_ms = int(round(float(t) * 1000))
    sign = "-" if total_ms < 0 else ""
    total_ms = abs(total_ms)
    h, rem = divmod(total_ms, 3_600_000)
    m, rem = divmod(rem, 60_000)
    s, ms = divmod(rem, 1000)
    return f"{sign}{h:02d}:{m:02d}:{s:02d}.{ms:03d}"


def utc_now():
    return datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds")


# ---------------------------------------------------------- frame index --------

class FrameIndex:
    """PTS of every frame of a video in presentation order (index == raw frame index)."""

    def __init__(self, pts, key, tb_num, tb_den, source_name="", source_size=0, source_mtime_ns=0):
        self.pts = np.asarray(pts, dtype=np.int64)
        self.key = np.asarray(key, dtype=bool)
        self.tb_num, self.tb_den = int(tb_num), int(tb_den)
        self.source_name = source_name
        self.source_size = int(source_size)
        self.source_mtime_ns = int(source_mtime_ns)
        self.n = len(self.pts)
        # video time: first frame is t = 0
        self.pts_s = (self.pts - self.pts[0]) * (self.tb_num / self.tb_den)
        self._key_idx = np.flatnonzero(self.key)

    @property
    def mean_gap_s(self):
        return float(self.pts_s[-1] / (self.n - 1)) if self.n > 1 else 0.0

    def nearest(self, t):
        """Frame index whose PTS is nearest to video time t, or None if t is outside the video."""
        gap = self.mean_gap_s
        if t < self.pts_s[0] - gap or t > self.pts_s[-1] + gap:
            return None
        i = int(np.searchsorted(self.pts_s, t))
        if i <= 0:
            return 0
        if i >= self.n:
            return self.n - 1
        return i - 1 if (t - self.pts_s[i - 1]) <= (self.pts_s[i] - t) else i

    def keyframe_at_or_before(self, k):
        j = int(np.searchsorted(self._key_idx, k, side="right")) - 1
        return int(self._key_idx[max(j, 0)])

    def summary(self):
        gaps = np.diff(self.pts_s) * 1000.0
        return {
            "frames": self.n, "duration_s": round(float(self.pts_s[-1]), 6),
            "mean_gap_ms": round(self.mean_gap_s * 1000.0, 6),
            "mean_fps": round(1.0 / self.mean_gap_s, 6) if self.mean_gap_s else None,
            "gap_min_ms": round(float(gaps.min()), 3) if len(gaps) else None,
            "gap_max_ms": round(float(gaps.max()), 3) if len(gaps) else None,
            "keyframes": int(self.key.sum()),
        }

    def save(self, csv_path):
        """Write the frame table (CSV) and its sidecar (.json with video identity and summary)."""
        csv_path = Path(csv_path)
        csv_path.parent.mkdir(parents=True, exist_ok=True)
        tmp = csv_path.with_name(csv_path.name + ".tmp")
        with open(tmp, "w", newline="") as f:
            f.write("frame,pts,time_s,time_hms,keyframe\n")
            f.writelines(f"{k},{p},{t:.6f},{fmt_time(t)},{int(key)}\n"
                         for k, (p, t, key) in enumerate(zip(self.pts.tolist(), self.pts_s.tolist(), self.key.tolist())))
        os.replace(tmp, csv_path)
        meta = {
            "pipeline_version": PIPELINE_VERSION, "created": utc_now(),
            "video_name": self.source_name, "video_size": self.source_size, "video_mtime_ns": self.source_mtime_ns,
            "time_base": f"{self.tb_num}/{self.tb_den}", "tb_num": self.tb_num, "tb_den": self.tb_den,
            "first_pts": int(self.pts[0]), "time_zero": "presentation time of the first frame",
            **self.summary(),
        }
        side = csv_path.with_suffix(".json")
        tmp = side.with_name(side.name + ".tmp")
        with open(tmp, "w") as f:
            json.dump(meta, f, indent=1)
        os.replace(tmp, side)

    @classmethod
    def load(cls, csv_path):
        csv_path = Path(csv_path)
        side = csv_path.with_suffix(".json")
        if not csv_path.exists() or not side.exists():
            raise SystemExit(f"Frame table not found ({csv_path.name} and its .json). Run index_video.py first.")
        with open(side) as f:
            meta = json.load(f)
        check_major_version(meta.get("pipeline_version"), side.name)
        arr = np.loadtxt(csv_path, delimiter=",", skiprows=1, usecols=(1, 4), dtype=np.int64)
        idx = cls(arr[:, 0], arr[:, 1].astype(bool), meta["tb_num"], meta["tb_den"],
                  meta["video_name"], meta["video_size"], meta["video_mtime_ns"])
        if idx.n != meta["frames"]:
            raise SystemExit(f"{csv_path.name} and {side.name} disagree on the frame count; re-run index_video.py")
        return idx

    def matches(self, video_path):
        """True when the video file has the size recorded when the frame table was made."""
        return os.stat(video_path).st_size == self.source_size


def _require_av():
    try:
        import av
    except ImportError:
        raise SystemExit("PyAV is required to read the video: pip install av")
    return av


def build_frame_index(video_path):
    """Demux the video (no decoding) and return its FrameIndex."""
    av = _require_av()
    video_path = Path(video_path)
    st = os.stat(video_path)
    pts, key = [], []
    with av.open(str(video_path)) as c:
        stream = c.streams.video[0]
        tb = stream.time_base
        for pkt in c.demux(stream):
            if pkt.pts is None:
                continue
            pts.append(pkt.pts)
            key.append(pkt.is_keyframe)
    if not pts:
        raise SystemExit(f"No video packets found in {video_path}")
    order = np.argsort(pts, kind="stable")
    pts = np.asarray(pts, dtype=np.int64)[order]
    key = np.asarray(key, dtype=bool)[order]
    if len(np.unique(pts)) != len(pts):
        raise SystemExit("Duplicate PTS values in the video; frame identity would be ambiguous")
    return FrameIndex(pts, key, tb.numerator, tb.denominator, video_path.name, st.st_size, st.st_mtime_ns)


def frame_table_path(force_dir, case_id):
    return Path(force_dir) / f"{case_id}_frame_times.csv"


def load_index(force_dir, case_id, video=None):
    """Load the frame table written by index_video.py; check it belongs to `video` when given."""
    idx = FrameIndex.load(frame_table_path(force_dir, case_id))
    if video is not None and not idx.matches(video):
        raise SystemExit(f"{Path(video).name} is not the video the frame table was made from "
                         "(different file size). Re-run index_video.py.")
    return idx


def half_window_frames_from_seconds(index, half_window_s):
    """Frames on each side of the centre frame for a half-window given in seconds (0.5 s = 15 at 30 fps)."""
    if half_window_s < 0:
        raise ValueError("half_window_s must be >= 0")
    return int(round(half_window_s / index.mean_gap_s)) if index.mean_gap_s else 0


# ------------------------------------------------------------ selection --------

@dataclass(frozen=True)
class SampleFrame:
    frame: int
    offset_frames: int
    offset_ms: float   # frame PTS minus the sample's target time, in ms


def max_half_window_frames(index, interval_s=1.0):
    """Largest half-window, in frames, that keeps every frame in one force second.

    Force samples are interval_s apart, so the centre frames of neighbouring force seconds are at least
    floor(interval_s / mean frame gap) frames apart (29 or 30 at 29.97 fps). The windows centre +/- h stay
    apart while 2h is less than that distance, which is strictly less than half the force-sample spacing."""
    if not index.mean_gap_s:
        return 0
    spacing = int(np.floor(interval_s / index.mean_gap_s + 1e-6))
    return max((spacing - 1) // 2, 0)


def check_half_window(index, half_window_frames):
    """Raise ValueError when a half-window would let one frame serve two force seconds."""
    limit = max_half_window_frames(index)
    if half_window_frames > limit:
        fps = 1.0 / index.mean_gap_s
        raise ValueError(
            f"a half-window of {half_window_frames} frames would let one frame serve two force seconds. "
            f"For this video ({fps:.3f} fps, {fps:.1f} frames between force samples) the largest allowed is "
            f"{limit} frames ({limit * index.mean_gap_s:.3f} s): strictly less than half the spacing. "
            "Run plan_measurements.py again (without --half-window-s it uses the largest allowed; a plan made by "
            "version 2.0.0 used 15). Marks already made are kept.")


def layout_offsets(half_window_frames, stride):
    """Frame offsets from the centre frame: 0, +/-stride, +/-2*stride, ... within +/-half_window_frames."""
    h, s = int(half_window_frames), int(stride)
    if h < 0 or s < 1:
        raise ValueError("half_window_frames must be >= 0 and stride >= 1")
    n = h // s
    return [j * s for j in range(-n, n + 1)]


def resolve_params(index=None, **overrides):
    """Sampling parameters: the defaults, then the given values. With a frame table the default
    half-window becomes the largest allowed for that video."""
    unknown = sorted(set(overrides) - set(PARAM_DEFAULTS))
    if unknown:
        raise ValueError(f"unknown sampling parameter(s): {', '.join(unknown)} "
                         f"(the 'edge'/'centred' layout choice was removed in version 2.1.0)")
    p = dict(PARAM_DEFAULTS)
    p.update({k: v for k, v in overrides.items() if v is not None})
    if index is not None and p["half_window_frames"] is None:
        p["half_window_frames"] = max_half_window_frames(index)
    return p


def target_time(force_second, params):
    return float(force_second) + params["time_offset_s"] + params["window_center_shift_s"]


def sample_frames(index, force_second, **params):
    """Return (centre_frame, [SampleFrame]) for one force second; (None, []) if outside the video."""
    p = resolve_params(index, **params)
    check_half_window(index, p["half_window_frames"])
    t = target_time(force_second, p)
    centre = index.nearest(t)
    if centre is None:
        return None, []
    out = []
    for off in layout_offsets(p["half_window_frames"], p["stride"]):
        k = centre + off
        if 0 <= k < index.n:
            out.append(SampleFrame(k, off, float((index.pts_s[k] - t) * 1000.0)))
    return centre, out


@dataclass
class PointPlan:
    samples: dict      # force second -> (centre_frame, [SampleFrame])
    frames: list       # sorted unique raw frame indices to measure
    membership: dict   # frame -> [(force second, offset_frames, offset_ms), ...]

    def shared_frames(self):
        """Frames that serve more than one force second. Empty unless force seconds are closer
        together than the half-window allows (for example a variable-frame-rate video with gaps)."""
        return sorted(k for k, m in self.membership.items() if len({x for x, _, _ in m}) > 1)


def plan_point(index, seconds, **params):
    """All frames to measure for one point (sorted, each once; see PointPlan.shared_frames)."""
    samples, membership = {}, {}
    for x in seconds:
        x = round(float(x), 3)
        centre, frames = sample_frames(index, x, **params)
        samples[x] = (centre, frames)
        for sf in frames:
            membership.setdefault(sf.frame, []).append((x, sf.offset_frames, sf.offset_ms))
    return PointPlan(samples, sorted(membership), membership)


# ---------------------------------------------------------- mark storage -------

class MarksStore:
    """Per-frame marks for one case, keyed by raw frame index. Every change is saved at once."""

    def __init__(self, path, case_id=""):
        self.path = Path(path)
        if self.path.exists():
            with open(self.path) as f:
                self.data = json.load(f)
            if self.data.get("schema_version") != SCHEMA_VERSION:
                raise SystemExit(f"{self.path.name}: unsupported schema_version {self.data.get('schema_version')}")
            check_major_version(self.data.get("pipeline_version"), self.path.name)
        else:
            self.data = {"schema_version": SCHEMA_VERSION, "pipeline_version": PIPELINE_VERSION,
                         "case_id": case_id, "video": {}, "points": {}}

    def set_video(self, index):
        self.data["video"] = {
            "name": index.source_name, "frames": index.n,
            "time_base": f"{index.tb_num}/{index.tb_den}", "first_pts": int(index.pts[0]),
        }

    def point(self, label):
        return self.data["points"].setdefault(label, {"params_history": [], "frames": {}})

    def note_params(self, label, params):
        """Append the parameter set when it differs from the last one used for this point."""
        hist = self.point(label)["params_history"]
        if not hist or any(hist[-1].get(k) != params[k] for k in PARAM_DEFAULTS):
            hist.append({**{k: params[k] for k in PARAM_DEFAULTS}, "at": utc_now()})

    def last_params(self, label):
        hist = self.data["points"].get(label, {}).get("params_history", [])
        return {k: hist[-1][k] for k in PARAM_DEFAULTS} if hist else None

    def get(self, label, frame):
        return self.data["points"].get(label, {}).get("frames", {}).get(str(int(frame)))

    def put(self, label, frame, rec):
        self.point(label)["frames"][str(int(frame))] = rec

    def delete(self, label, frame):
        self.point(label)["frames"].pop(str(int(frame)), None)

    # -- camera check: scale of each frame's scene relative to the point's calibration frame
    def set_scale(self, label, frame, scale, inliers):
        """Store the estimated scale (None when no estimate could be made) and the number of matched features."""
        self.point(label).setdefault("scale", {})[str(int(frame))] = {
            "scale": None if scale is None else round(float(scale), 4), "inliers": int(inliers)}

    def get_scale(self, label, frame):
        return self.data["points"].get(label, {}).get("scale", {}).get(str(int(frame)))

    def set_scale_check(self, label, **meta):
        """What the camera check was run against (calibration frame and image, thresholds, time)."""
        self.point(label)["scale_check"] = {**meta, "checked_at": utc_now(), "pipeline_version": PIPELINE_VERSION}

    def get_scale_check(self, label):
        return self.data["points"].get(label, {}).get("scale_check")

    def counts(self, label, frames):
        marked = skipped = 0
        for k in frames:
            r = self.get(label, k)
            if r is None:
                continue
            if r.get("skipped"):
                skipped += 1
            else:
                marked += 1
        return marked, skipped, len(frames) - marked - skipped

    def save(self):
        """Atomic save: write a temp file, keep the previous version as .bak, then replace."""
        self.data["pipeline_version"] = PIPELINE_VERSION
        tmp = self.path.with_name(self.path.name + ".tmp")
        with open(tmp, "w") as f:
            json.dump(self.data, f, indent=1)
            f.flush()
            os.fsync(f.fileno())
        if self.path.exists():
            os.replace(self.path, self.path.with_name(self.path.name + ".bak"))
        os.replace(tmp, self.path)


# ------------------------------------------------- selected force rows ---------

def selected_force_rows(workbook_path, label):
    """[(force second, role)] for every row of a point's sheet that is to be measured, plus a legacy flag.

    Roles (from extract_point_forces.py): window | buffer-before | buffer-after. A workbook made without a
    Role column (older extraction) falls back to its 'In Window' rows, all counted as 'window'.
    """
    import openpyxl
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = f"Point {label}"
        if sheet not in wb.sheetnames:
            raise SystemExit(f"Sheet '{sheet}' not found in {Path(workbook_path).name} (sheets: {wb.sheetnames})")
        rows = wb[sheet].iter_rows(values_only=True)
        headers = list(next(rows))
        sec = headers.index("Seconds")
        if "Role" in headers:
            role = headers.index("Role")
            out = [(round(float(r[sec]), 3), r[role]) for r in rows if r[sec] is not None and r[role] in ROLES]
            legacy = False
        else:
            win = headers.index("In Window")
            out = [(round(float(r[sec]), 3), "window") for r in rows if r[sec] is not None and r[win] == "YES"]
            legacy = True
    finally:
        wb.close()
    return out, legacy


# ------------------------------------------------------- point segments -------

def parse_hms(text):
    """'HH:MM:SS', 'HH:MM:SS.mmm' or plain seconds -> seconds. ValueError for anything else."""
    parts = str(text).strip().split(":")
    try:
        if len(parts) == 1:
            return float(parts[0])
        if len(parts) == 3:
            return int(parts[0]) * 3600 + int(parts[1]) * 60 + float(parts[2])
    except ValueError:
        pass
    raise ValueError(f"not a time (use HH:MM:SS): {text!r}")


def read_selection(workbook_path):
    """{point: {start_s, end_s, start_text, end_text, buffer_s, time_offset_s}} from the Selection sheet that
    extract_point_forces.py writes (video times, as entered). {} when the workbook has no such sheet."""
    import openpyxl
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "Selection" not in wb.sheetnames:
            return {}
        rows = wb["Selection"].iter_rows(values_only=True)
        headers = list(next(rows))
        col = {h: i for i, h in enumerate(headers)}
        out = {}
        for r in rows:
            if r[0] not in ("A", "B", "C", "D"):
                continue
            out[r[0]] = {"start_s": parse_hms(r[col["Start (as entered)"]]),
                         "end_s": parse_hms(r[col["End (as entered)"]]),
                         "start_text": str(r[col["Start (as entered)"]]), "end_text": str(r[col["End (as entered)"]]),
                         "buffer_s": float(r[col["Buffer (s)"]]), "time_offset_s": float(r[col["Time offset (s)"]])}
        return out
    finally:
        wb.close()


def segment_position(seg, t):
    """Where video time t lies in a point's segment: window | buffer-before | buffer-after | outside the selected range."""
    start, end, buf = seg["start_s"], seg["end_s"], seg["buffer_s"]
    if start <= t <= end:
        return "window"
    if start - buf <= t < start:
        return "buffer-before"
    if end < t <= end + buf:
        return "buffer-after"
    return "outside the selected range"


def first_frame_at_or_after(index, t):
    """Index of the first frame whose stored time is at or after t (a microsecond of slack), or index.n if none."""
    return int(np.searchsorted(index.pts_s, t - 1e-6))


def last_frame_at_or_before(index, t):
    """Index of the last frame whose stored time is at or before t (a microsecond of slack), or -1 if none."""
    return int(np.searchsorted(index.pts_s, t + 1e-6, side="right")) - 1


def segment_frame_range(index, seg):
    """(first, last) frame of a point's window plus buffer: the first frame at or after the window start minus the
    buffer and the last frame at or before the window end plus the buffer, so both lie inside the selected range."""
    return (first_frame_at_or_after(index, seg["start_s"] - seg["buffer_s"]),
            last_frame_at_or_before(index, seg["end_s"] + seg["buffer_s"]))


def resolve_image(path, base):
    """An image path as stored in a calibration JSON (absolute, or relative to the repository folder or to the
    working directory), or None if no such file exists."""
    p = Path(path)
    for candidate in ([p] if p.is_absolute() else [Path(base) / p, Path.cwd() / p]):
        if candidate.exists():
            return str(candidate)
    return None


def previous_reference(case_id, label, base):
    """The frame to show beside the one being calibrated: the previous point's (A shows the ruler frame, B shows
    A's, C shows B's, D shows C's), else the ruler frame when the previous point has no calibration yet.

    Returns {name, image (path or None), p1, p2, missing (stored path when the image is gone)} or None."""
    fdir = Path(base) / "Videos" / case_id / "force_data"
    candidates = []
    if label in ("B", "C", "D"):
        prev = "ABCD"["ABCD".index(label) - 1]
        candidates.append((f"point {prev}", fdir / f"{case_id}_{prev}_calibration.json"))
    candidates.append(("the ruler frame", fdir / f"{case_id}_calibration.json"))
    for name, path in candidates:
        if not path.exists():
            continue
        try:
            cal = json.loads(path.read_text())
            p1, p2 = cal["landmark_coords_p1"], cal["landmark_coords_p2"]
            stored = cal["calibration_frame"]
        except (OSError, ValueError, KeyError):
            continue
        image = resolve_image(stored, base)
        return {"name": name, "image": image, "p1": tuple(p1), "p2": tuple(p2),
                "missing": None if image else str(stored)}
    return None


def reference_crop(shape, p1, p2, factor=3.0, floor=300):
    """(x0, x1, y0, y1): a square of the image centred on the midpoint of two landmark clicks, its side `factor` times
    their distance (at least `floor` px), shifted to stay inside the image and never larger than its short side."""
    h, w = shape[:2]
    dist = float(np.hypot(p2[0] - p1[0], p2[1] - p1[1]))
    side = int(min(round(max(factor * dist, floor)), h, w))
    cx, cy = (p1[0] + p2[0]) / 2.0, (p1[1] + p2[1]) / 2.0
    x0 = int(min(max(round(cx - side / 2.0), 0), w - side))
    y0 = int(min(max(round(cy - side / 2.0), 0), h - side))
    return x0, x0 + side, y0, y0 + side


def load_rgb(path):
    """An image file as an RGB uint8 array (Pillow, so no OpenCV is needed)."""
    from PIL import Image
    with Image.open(path) as im:
        return np.asarray(im.convert("RGB"))


def draw_landmark_crop(fig, rect, rgb, p1, p2, title):
    """A reference panel: the square around two landmark clicks, enlarged, with the clicks numbered in order, and a small
    thumbnail of the whole frame with the square outlined. rect = (left, bottom, width, height) in figure fractions.
    Returns the two axes."""
    from matplotlib.patches import Rectangle
    left, bottom, width, height = rect
    x0, x1, y0, y1 = reference_crop(rgb.shape, p1, p2)
    ax = fig.add_axes(rect)
    ax.imshow(rgb, interpolation="nearest")
    ax.set_xlim(x0, x1)
    ax.set_ylim(y1, y0)
    ax.axis("off")
    ax.set_title(title, fontsize=9)
    ax.plot([p1[0], p2[0]], [p1[1], p2[1]], "-", color="cyan", linewidth=1.2)
    for n, (x, y) in enumerate((p1, p2), 1):
        ax.plot(x, y, "+", color="cyan", markersize=14, markeredgewidth=1.6)
        ax.annotate(str(n), (x, y), color="yellow", fontsize=12, fontweight="bold",
                    xytext=(7, -7), textcoords="offset points")
    thumb = fig.add_axes([left + 0.005, bottom + 0.005, width * 0.36, height * 0.30])
    thumb.imshow(rgb)
    thumb.add_patch(Rectangle((x0, y0), x1 - x0, y1 - y0, fill=False, edgecolor="yellow", linewidth=1.2))
    thumb.axis("off")
    return ax, thumb


def save_frame_png(rgb, path):
    """Write an RGB frame as a lossless PNG at its own resolution (Pillow, which matplotlib already needs)."""
    from PIL import Image
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Image.fromarray(np.ascontiguousarray(rgb)).save(path)


# ------------------------------------------------------ measurement plan -------

def plan_path(force_dir, case_id):
    return Path(force_dir) / f"{case_id}_measurement_plan.json"


def load_plan(force_dir, case_id, check_version=True):
    """The measurement plan written by plan_measurements.py, or None if there is none. A plan from another major
    version is refused unless check_version is False (plan_measurements.py reads the old one only to say it replaces it)."""
    path = plan_path(force_dir, case_id)
    if not path.exists():
        return None
    with open(path) as f:
        plan = json.load(f)
    if plan.get("schema_version") != PLAN_SCHEMA_VERSION:
        raise SystemExit(f"{path.name}: unsupported schema_version {plan.get('schema_version')}")
    if check_version:
        check_major_version(plan.get("pipeline_version"), path.name)
    return plan


def plan_point_params(plan, label):
    """Sampling parameters for one point from a measurement plan."""
    p = plan["parameters"]
    pt = plan["points"].get(label)
    if pt is None:
        raise SystemExit(f"Point {label} is not in the measurement plan; re-run plan_measurements.py")
    return {"half_window_frames": p["half_window_frames"], "stride": pt["stride"],
            "time_offset_s": p["time_offset_s"], "window_center_shift_s": p["window_center_shift_s"]}


# -------------------------------------------------------------- decoding -------

def frame_crc32(rgb):
    return zlib.crc32(np.ascontiguousarray(rgb)) & 0xFFFFFFFF


class FrameReader:
    """Exact-frame decoder: returns the RGB array of raw frame k and checks its PTS."""

    CACHE_FRAMES = 24

    def __init__(self, video_path, index):
        av = _require_av()
        self.index = index
        self._container = av.open(str(video_path))
        self._stream = self._container.streams.video[0]
        self._stream.thread_type = "AUTO"
        self._iter = None
        self._last = None          # raw frame index of the last frame the decoder yielded
        self._cache = OrderedDict()

    def close(self):
        self._container.close()

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        self.close()

    def _seek(self, target_pts):
        self._container.seek(int(target_pts), stream=self._stream, backward=True, any_frame=False)
        self._stream.codec_context.flush_buffers()
        self._iter = self._container.decode(self._stream)
        self._last = None

    def get(self, k):
        k = int(k)
        if k in self._cache:
            self._cache.move_to_end(k)
            return self._cache[k]
        target = int(self.index.pts[k])
        seek_cost = k - self.index.keyframe_at_or_before(k) + 1
        fwd_cost = (k - self._last) if (self._last is not None and k > self._last) else None
        if self._iter is None or fwd_cost is None or fwd_cost > seek_cost + 30:
            self._seek(target)
        for frame in self._iter:
            if frame.pts == target:
                self._last = k
                rgb = frame.to_ndarray(format="rgb24")
                self._cache[k] = rgb
                while len(self._cache) > self.CACHE_FRAMES:
                    self._cache.popitem(last=False)
                return rgb
            if frame.pts is not None and frame.pts > target:
                break
        raise RuntimeError(f"Could not decode frame {k} (PTS {target}); the video may have a gap or odd timestamps")
