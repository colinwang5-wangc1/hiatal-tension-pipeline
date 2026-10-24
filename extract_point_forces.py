"""
extract_point_forces.py — Select the force samples to measure for points A-D from the patient Excel sheet.

For each point, takes the start and end times noted while scrubbing the video (HH:MM:SS; fractions are
accepted), adds a buffer of BUFFER_SECS before and after, and keeps every force sample in that range.
Every kept row is measured later, and each row is labelled with its role:
    window         between the noted start and end
    buffer-before  in the buffer before the start
    buffer-after   in the buffer after the end
so buffer rows can be kept as data but left out of an analysis (fit_moduli.py --rows window).

Times in the sheet are VIDEO time (what a media player shows). A force sample x is at video time
x + time_offset_s (video time minus force time; 0 unless the two clocks are known to differ).

Saves one Excel file per case with one sheet per point plus a "Selection" sheet that records what was
selected: the start/end as entered, the buffer, the offset and the row counts.

Usage:
    python3 extract_point_forces.py [case_id] [--buffer-s S] [--time-offset-s S] [--tracking-sheet FILE]
(without a case_id every case in the sheet is processed; the sheet defaults to
"Hiatal tension patients.xlsx" in the repository root)
"""

import argparse
import json
import openpyxl
from openpyxl.styles import PatternFill, Font
from pathlib import Path
import sys

import framelib as fl

# How many seconds before/after each window to include
BUFFER_SECS = 5

BASE = Path(__file__).resolve().parent
PATIENT_FILE = BASE / "Hiatal tension patients.xlsx"


def ts_to_secs(ts):
    """Convert HH:MM:SS.mmm string to float seconds."""
    if not ts:
        return None
    parts = str(ts).split(":")
    h, m, s = int(parts[0]), int(parts[1]), float(parts[2])
    return h * 3600 + m * 60 + s


def secs_to_ts(secs):
    """Convert float seconds to HH:MM:SS.mmm string."""
    h = int(secs // 3600)
    m = int((secs % 3600) // 60)
    s = secs % 60
    return f"{h:02d}:{m:02d}:{s:06.3f}"


def load_force_arm(json_path, arm):
    """Return list of (x_secs, y_force) for a given arm."""
    with open(json_path) as f:
        data = json.load(f)
    arm_data = data.get(arm, [])
    return [(pt["x"], pt["y"]) for pt in arm_data if pt.get("y") is not None]


def select_rows(force_seconds, start_s, end_s, buffer_s=BUFFER_SECS, time_offset_s=0.0):
    """[(force second, role)] for the force samples to measure.

    start_s and end_s are video times; a force sample x sits at video time x + time_offset_s.
    Samples from start - buffer to end + buffer (inclusive) are kept.
    """
    out = []
    for x in sorted(set(force_seconds)):
        t = x + time_offset_s
        if start_s - buffer_s <= t < start_s:
            role = "buffer-before"
        elif start_s <= t <= end_s:
            role = "window"
        elif end_s < t <= end_s + buffer_s:
            role = "buffer-after"
        else:
            continue
        out.append((x, role))
    return out


def leftover_force(rows, arm_lookup):
    """'USM1 1.7 N, USM3 0.4 N' for the arms with force on the first selected row, or '' if there is none."""
    if not rows:
        return ""
    x = rows[0][0]
    found = [(arm, vals.get(x)) for arm, vals in arm_lookup.items()]
    return ", ".join(f"{arm} {v:g} N" for arm, v in found if v is not None and abs(v) > 1e-6)


def parse_args(argv=None):
    ap = argparse.ArgumentParser(description="Select force samples for the measurement windows of each point.")
    ap.add_argument("case_id", nargs="?", help="process only this case (default: every case in the sheet)")
    ap.add_argument("--buffer-s", type=float, default=BUFFER_SECS, dest="buffer_s",
                    help=f"seconds of buffer before and after each window [{BUFFER_SECS}]")
    ap.add_argument("--time-offset-s", type=float, default=0.0, dest="time_offset_s",
                    help="video time minus force time [0]")
    ap.add_argument("--tracking-sheet", default=None, dest="tracking_sheet",
                    help="spreadsheet with the Point A-D start and end times [Hiatal tension patients.xlsx]")
    ap.add_argument("--version", action="version", version=f"%(prog)s {fl.PIPELINE_VERSION}")
    return ap.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    if args.buffer_s < 0:
        sys.exit("ERROR: --buffer-s must be >= 0")

    # Load patient sheet
    sheet_path = Path(args.tracking_sheet) if args.tracking_sheet else PATIENT_FILE
    if not sheet_path.exists():
        sys.exit(f"ERROR: tracking spreadsheet not found: {sheet_path}")
    wb_in = openpyxl.load_workbook(sheet_path)
    ws = wb_in.active

    headers = [cell.value for cell in ws[1]]
    col = {h: i for i, h in enumerate(headers)}
    required = ["Study ID"] + [f"Point {p} {side}" for p in "ABCD" for side in ("start", "end")]
    missing = [h for h in required if h not in col]
    if missing:
        sys.exit(f"ERROR: {sheet_path.name} is missing required column(s): {', '.join(missing)}.\n"
                 "       The first row must contain: " + ", ".join(required) + " (leave a point's cells empty if it was not measured).")

    seen, written, skipped = False, [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        case = row[col["Study ID"]]
        if not case:
            continue
        if args.case_id and str(case) != args.case_id:
            continue
        seen = True

        # Check if any points are defined
        points = {}
        entered = {}
        for pt in ["A", "B", "C", "D"]:
            start = row[col[f"Point {pt} start"]]
            end   = row[col[f"Point {pt} end"]]
            if start and end:
                points[pt] = (ts_to_secs(str(start)), ts_to_secs(str(end)))
                entered[pt] = (str(start), str(end))

        if not points:
            print(f"{case}: no points defined, skipping")
            skipped.append(f"{case}: no points defined (fill in the Point A-D start and end times)")
            continue

        json_path = BASE / "Videos" / case / "force_data" / f"{case}_force_raw.json"
        if not json_path.exists():
            print(f"{case}: force JSON not found, skipping")
            skipped.append(f"{case}: force JSON not found ({json_path.name}; run parse_force_api.py or put the "
                           f"console export there)")
            continue

        # Determine active arms
        with open(json_path) as f:
            raw = json.load(f)
        arms = [arm for arm in ["USM1", "USM2", "USM3", "USM4"]
                if arm in raw and len(raw[arm]) > 0]

        print(f"\n{case}: points {list(points.keys())}, arms {arms}, buffer {args.buffer_s:g} s, "
              f"time offset {args.time_offset_s:g} s")

        # Load force data per arm
        arm_data = {arm: load_force_arm(json_path, arm) for arm in arms}

        # Build output workbook
        wb_out = openpyxl.Workbook()
        wb_out.remove(wb_out.active)  # remove default sheet

        highlight = PatternFill(start_color="FFFFD700", end_color="FFFFD700", fill_type="solid")
        buffer_fill = PatternFill(start_color="FFDDEBF7", end_color="FFDDEBF7", fill_type="solid")
        bold = Font(bold=True)
        selection = []

        for pt_name, (start_s, end_s) in points.items():
            ws_out = wb_out.create_sheet(title=f"Point {pt_name}")

            # Header row
            header = ["Timestamp", "Seconds"] + arms + ["In Window", "Role"]
            for c, h in enumerate(header, 1):
                cell = ws_out.cell(row=1, column=c, value=h)
                cell.font = bold

            # Every force second seen on any arm, then the ones inside the buffered range
            all_x = {x for arm in arms for x, _ in arm_data[arm]}
            rows = select_rows(all_x, start_s, end_s, args.buffer_s, args.time_offset_s)

            # Build lookup dicts
            arm_lookup = {}
            for arm in arms:
                arm_lookup[arm] = {x: y for x, y in arm_data[arm]}

            # Write rows
            for r, (x, role) in enumerate(rows, 2):
                ws_out.cell(row=r, column=1, value=secs_to_ts(x))
                ws_out.cell(row=r, column=2, value=round(x, 3))
                for c, arm in enumerate(arms, 3):
                    ws_out.cell(row=r, column=c, value=arm_lookup[arm].get(x))
                ws_out.cell(row=r, column=len(arms) + 3, value="YES" if role == "window" else "")
                ws_out.cell(row=r, column=len(arms) + 4, value=role)

                # Window rows gold, buffer rows light blue
                fill = highlight if role == "window" else buffer_fill
                for c in range(1, len(arms) + 5):
                    ws_out.cell(row=r, column=c).fill = fill

            # Column widths
            ws_out.column_dimensions["A"].width = 16
            for i in range(2, len(arms) + 5):
                ws_out.column_dimensions[chr(64 + i)].width = 14

            counts = {role: sum(1 for _, rl in rows if rl == role) for role in fl.ROLES}
            xs = [x for x, _ in rows]
            selection.append([
                pt_name, entered[pt_name][0], entered[pt_name][1], args.buffer_s, args.time_offset_s,
                counts["window"], counts["buffer-before"], counts["buffer-after"],
                f"{min(xs):g} to {max(xs):g}" if xs else "none"])
            print(f"  Point {pt_name}: {start_s:.1f}s–{end_s:.1f}s → {len(rows)} rows "
                  f"({counts['window']} window, {counts['buffer-before']} buffer before, "
                  f"{counts['buffer-after']} buffer after)")
            leftover = leftover_force(rows, arm_lookup)
            if leftover:
                print(f"  WARNING: point {pt_name}: force is already nonzero on the first row ({secs_to_ts(rows[0][0])}): "
                      f"{leftover}. This is force left over from earlier activity at the start of the buffer. "
                      "Convention B takes its reference after the window start, so this alone does not make it "
                      "undefined. Note that a blank arm reading counts as 0 N in the totals and the fits, so a "
                      "blank row can look like a zero-force row when an arm was simply not logged.")

        # The record of what was selected
        ws_sel = wb_out.create_sheet("Selection")
        ws_sel.append(["Point", "Start (as entered)", "End (as entered)", "Buffer (s)", "Time offset (s)",
                       "Window rows", "Buffer-before rows", "Buffer-after rows", "Force seconds kept"])
        for c in ws_sel[1]:
            c.font = bold
        for line in selection:
            ws_sel.append(line)
        ws_sel.append([])
        ws_sel.append(["Start and end are video time (what a media player shows). A force sample x is at video time "
                       "x + time offset. Rows from start - buffer to end + buffer are kept; every kept row is measured."])
        ws_sel.append([f"Pipeline version {fl.PIPELINE_VERSION}"])
        for letter in "ABCDEFGHI":
            ws_sel.column_dimensions[letter].width = 18

        out_path = BASE / "Videos" / case / "force_data" / f"{case}_points_ABCD_force.xlsx"
        wb_out.save(out_path)
        written.append(case)
        print(f"  Saved: {out_path}")

    if args.case_id and not seen:
        sys.exit(f"ERROR: case {args.case_id!r} is not in the tracking sheet {sheet_path.name} (Study ID column)")
    if not written:
        why = "; ".join(skipped) if skipped else "the tracking sheet has no cases"
        sys.exit(f"ERROR: no output was written ({why})")


if __name__ == "__main__":
    main()
